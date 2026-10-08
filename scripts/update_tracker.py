"""One-step tracker update, no Claude needed:
moves Leads rows marked Add into the Tracker (academic and, if set up, industry tracker),
refreshes the shared letter-writer file, and copies updated master files from My Materials
into the job market folder (sync_materials.py; skipped when no materials index exists).

  python3 scripts/update_tracker.py

If Excel has the tracker open, it asks Excel to save and close it, updates it, and opens it
again (macOS asks once to allow controlling Excel). It also puts a clickable
"Update Tracker" link at the top of the Tracker and Leads sheets that runs the macOS Shortcuts
shortcut named "Update Tracker" (a "Run Shell Script" action calling this script).
Apply Via and Letters? are guessed from the ad; ask Claude to "fill the tracker" for a full check.
"""
import os
import subprocess
import sys
import urllib.parse

from openpyxl import load_workbook
from openpyxl.styles import Font

from common import close_in_excel, load_config, reopen_in_excel, restore_dropdowns, trackers
from leads import promote

HERE = os.path.dirname(os.path.abspath(__file__))
LABEL = "▶ Update Tracker (click: moves Leads marked Add, refreshes the letter list)"
SHORTCUT = "Update Tracker"   # macOS Shortcuts shortcut that runs this script (Excel may open its URL)


def add_buttons(path):
    """Excel's sandbox will not launch a .command file, but it may open a shortcuts:// URL."""
    wb = load_workbook(path)
    url = "shortcuts://run-shortcut?name=" + urllib.parse.quote(SHORTCUT)
    for name, cell in (("Tracker", "H1"), ("Leads", "H1")):
        if name not in wb.sheetnames:
            continue
        c = wb[name][cell]
        if c.value != LABEL or c.hyperlink is None or c.hyperlink.target != url:
            c.value = LABEL
            c.hyperlink = url
            c.font = Font(name="Arial", size=11, bold=True, color="0563C1", underline="single")
    restore_dropdowns(wb, load_config()["letter_writers"])
    wb.save(path)


cfg = load_config()
moved, reopen = [], []
for kind, tracker in trackers(cfg):
    reopen.append((tracker, close_in_excel(tracker)))
    moved += promote(cfg, tracker, kind)
    add_buttons(tracker)
subprocess.run([sys.executable, os.path.join(HERE, "export_letter_list.py")], check=True)
if cfg.get("materials_dir"):
    subprocess.run([sys.executable, os.path.join(HERE, "sync_materials.py")])
print(f"Done: {len(moved)} new row(s) in the Tracker(s); letter-writer file refreshed.")
for tracker, was_open in reopen:
    reopen_in_excel(tracker, was_open)
