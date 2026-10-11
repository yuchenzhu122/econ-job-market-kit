#!/usr/bin/env python3
"""Sort the Tracker sheet by deadline, then fit, then status.

    python3 sort_tracker.py [--industry] [--dry-run] [--prune-skip]

Rows with no deadline go after dated ones (by fit and status); empty rows stay at the bottom; the
industry tracker's EXAMPLE row stays first. Adds a Fit column (High / Medium / Low) when missing,
filled from the "From scan (X fit)" note that leads.py promote writes. --prune-skip (industry tracker): first delete rows whose Verdict is "skip" and mark their Leads
row Decision = Pass, so they are not promoted again. Formulas (Days Left,
_sort, _order) are rewritten from the row-5 template after the move. Also widens the header
filter to every column so you can re-sort in Excel.
"""
import argparse
import re
from copy import copy

import openpyxl
from openpyxl.utils import get_column_letter as L

from common import (close_in_excel, ensure_text_column, load_config, reopen_in_excel,
                    restore_dropdowns, tracker_headers, tracker_path)

FIT = {"High": 0, "Medium": 1, "Low": 2}
NOTE_FIT = re.compile(r"From scan \((High|Medium|Low) fit\)")


def at(formula, r):
    """Point a row-5 template formula at row r (relative row-5 refs only)."""
    return re.sub(r"(\$?)([A-Z]{1,3})(\$?)(\d+)",
                  lambda m: m.group(0) if (m.group(3) or m.group(4) != "5") else f"{m.group(1)}{m.group(2)}{r}",
                  formula)


def prune_skip(wb):
    """Delete Tracker rows with Verdict = skip; set the matching Leads row's Decision to Pass.
    Returns the (employer, position) pairs removed."""
    ws = wb["Tracker"]
    H = tracker_headers(ws)
    if "Verdict" not in H:
        return []
    gone = []
    for r in range(5, ws.max_row + 1):
        if str(ws.cell(row=r, column=H["Verdict"]).value or "").strip().lower() != "skip":
            continue
        if not ws.cell(row=r, column=H["Employer"]).value:
            continue
        gone.append((ws.cell(row=r, column=H["Employer"]).value, ws.cell(row=r, column=H["Position"]).value,
                     ws.cell(row=r, column=H["Link"]).value))
        for k in range(1, ws.max_column + 1):
            c = ws.cell(row=r, column=k)
            if not (isinstance(c.value, str) and c.value.startswith("=")):
                c.value = None
            c.hyperlink = None
    if gone and "Leads" in wb.sheetnames:
        li = wb["Leads"]
        LH = {li.cell(row=4, column=c).value: c for c in range(1, li.max_column + 1)}
        for r in range(5, li.max_row + 1):
            e, p, l = (li.cell(row=r, column=LH[h]).value for h in ("Employer", "Position", "Link"))
            if any(e == ge and (l == gl or p == gp) for ge, gp, gl in gone):
                li.cell(row=r, column=LH["Decision"], value="Pass")
    return [(e, p) for e, p, _ in gone]


def sort_sheet(wb, dry_run=False):
    """Sort wb's Tracker sheet in place (caller saves). Returns the number of rows sorted."""
    ws = wb["Tracker"]
    H = tracker_headers(ws)
    if "Fit" not in H:
        H["Fit"] = ensure_text_column(ws, "Fit", width=9)
    notes, fitc = H["Notes"], H["Fit"]
    status_order = []
    li = wb["Lists"]
    for c in range(1, li.max_column + 1):
        if li.cell(row=3, column=c).value == "Status":
            status_order = [li.cell(row=r, column=c).value for r in range(4, 31) if li.cell(row=r, column=c).value]
    ncol = max(H.values())
    last = max(154, ws.max_row)
    formula_cols = [c for c in range(1, ncol + 1)
                    if isinstance(ws.cell(row=5, column=c).value, str) and ws.cell(row=5, column=c).value.startswith("=")]
    tmpl = {c: ws.cell(row=5, column=c).value for c in formula_cols}

    rows = []
    for r in range(5, last + 1):
        if not ws.cell(row=r, column=H["Employer"]).value:
            continue
        if not ws.cell(row=r, column=fitc).value:
            m = NOTE_FIT.search(str(ws.cell(row=r, column=notes).value or ""))
            if m:
                ws.cell(row=r, column=fitc, value=m.group(1))
        cells = [(c.value, copy(c._style), copy(c.hyperlink) if c.hyperlink else None)
                 for c in (ws.cell(row=r, column=k) for k in range(1, ncol + 1))]
        rows.append((r, cells, ws.row_dimensions[r].height))

    def key(item):
        r = item[0]
        emp = str(ws.cell(row=r, column=H["Employer"]).value)
        dl = ws.cell(row=r, column=H["Deadline"]).value
        st = ws.cell(row=r, column=H["Status"]).value
        return (not emp.startswith("EXAMPLE"),
                dl is None, dl.toordinal() if hasattr(dl, "toordinal") else 0,
                FIT.get(ws.cell(row=r, column=fitc).value, 3),
                status_order.index(st) if st in status_order else len(status_order))

    rows.sort(key=key)
    if dry_run:
        for r, cells, _ in rows:
            v = {h: cells[c - 1][0] for h, c in H.items()}
            print(r, v["Deadline"], v["Fit"], v["Status"], v["Employer"], sep=" | ")
        return len(rows)
    for r in range(5, last + 1):  # clear the data block, then write rows back in order
        for k in range(1, ncol + 1):
            c = ws.cell(row=r, column=k)
            c.hyperlink = None
            if k not in formula_cols:
                c.value = None
    for i, (_, cells, height) in enumerate(rows):
        r = 5 + i
        for k, (v, style, link) in enumerate(cells, start=1):
            c = ws.cell(row=r, column=k)
            c.value, c._style = v, style
            if link:
                link.ref = c.coordinate
                c.hyperlink = link
        ws.row_dimensions[r].height = height
    for r in range(5, 5 + len(rows)):  # the Link cell's text is the source of truth for its hyperlink
        c = ws.cell(row=r, column=H["Link"])
        c.hyperlink = c.value if str(c.value or "").startswith("http") else None
    for r in range(5, last + 1):
        for c in formula_cols:
            ws.cell(row=r, column=c, value=at(tmpl[c], r))
    ws.auto_filter.ref = f"A4:{L(ncol)}{last}"
    return len(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--industry", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--prune-skip", action="store_true", help="delete Verdict = skip rows first")
    a = ap.parse_args()
    cfg = load_config()
    path = tracker_path(cfg, "industry" if a.industry else "academic")
    was = close_in_excel(path) if not a.dry_run else False
    wb = openpyxl.load_workbook(path)
    if a.prune_skip and not a.dry_run:
        for e, p in prune_skip(wb):
            print(f"removed (skip): {e} | {p}")
    n = sort_sheet(wb, a.dry_run)
    if a.dry_run:
        return
    restore_dropdowns(wb, cfg["letter_writers"])
    wb.save(path)
    reopen_in_excel(path, was)
    print(f"Sorted {n} rows in {path}")


if __name__ == "__main__":
    main()
