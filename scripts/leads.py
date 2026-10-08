"""Manage the Leads sheet of the tracker.

  python3 scripts/leads.py add <evaluated.json>   add evaluated postings to the top of the Leads sheet
  python3 scripts/leads.py promote                 copy Leads rows with Decision = Add into Tracker
  python3 scripts/leads.py sort                    re-sort existing Leads rows, newest Found date first
Add --industry to work on the industry tracker instead of the academic one.

Before a posting is written to Leads (or promoted), every tracker is checked for the same job (same
link, or same employer and a near-identical title). A job already in a Tracker sheet is not added
again; one only in the other tracker's Leads is added with a flag.

evaluated.json is a list of objects with keys:
  id, url, source, employer, position, track (industry: category), type, location, deadline (YYYY-MM-DD or ""),
  fit (High | Medium | Low), why, flags (string or list), also_at (optional list of other links)
deadline may also be text like "Review begins 2026-10-30": the date is used and the text kept in Flags.

Leads columns: Found, Fit, Employer, Position, Track (industry: Category), Type, Location, Deadline, Why, Flags,
Source, Link, Decision (Add / Maybe / Pass; "Added" once promoted).
"""
import datetime as dt
import json
import os
import re
import sys
from copy import copy

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from common import (close_in_excel, ensure_date_column, jm_path, load_config, norm_link, reopen_in_excel,
                    restore_dropdowns, tracker_path, trackers)

COLS = [("Found", 11), ("Fit", 8), ("Employer", 26), ("Position", 30), ("Track", 12), ("Type", 15),
        ("Location", 16), ("Deadline", 11), ("Why", 44), ("Flags", 26), ("Source", 7), ("Link", 22),
        ("Decision", 10)]
FONT = "Arial"
side = Side(style="thin", color="D0D0D0")
BD = Border(left=side, right=side, top=side, bottom=side)
FIT_ORDER = {"High": 0, "Medium": 1, "Low": 2}


def cols(kind):
    return [("Category", w) if h == "Track" and kind == "industry" else (h, w) for h, w in COLS]


def ensure_sheet(wb, accent, kind="academic"):
    if "Leads" in wb.sheetnames:
        return wb["Leads"]
    ws = wb.create_sheet("Leads", 2)
    ws["A1"] = ("Leads · found in your LinkedIn / Indeed job alerts and on company career pages" if kind == "industry"
                else "Leads · found automatically on JOE, EconJobMarket, Chronicle and Inside Higher Ed")
    ws["A1"].font = Font(name=FONT, size=15, bold=True, color=accent)
    ws["A2"] = ("Pick a Decision for each row: Add = move it to the Tracker (ask Claude 'add my leads to the "
                "tracker'), Maybe = keep, Pass = ignore. Fit is Claude's judgment against your CV; check the Why and Flags.")
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="555555")
    for i, (h, w) in enumerate(cols(kind), start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=accent)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BD
        ws.column_dimensions[L(i)].width = w
    dv = DataValidation(type="list", formula1='"Add,Maybe,Pass,Added"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"M5:M2000")
    for val, fill, color in (("High", "C6EFCE", "006100"), ("Medium", "FFEB9C", "9C5700"), ("Low", "EDEDED", "7F7F7F")):
        ws.conditional_formatting.add("B5:B2000", FormulaRule(formula=[f'$B5="{val}"'],
                                      fill=PatternFill("solid", fgColor=fill), font=Font(name=FONT, color=color, bold=True)))
    for val, fill in (("Add", "BDD7EE"), ("Added", "D9D9D9"), ("Pass", "F2F2F2")):
        ws.conditional_formatting.add("M5:M2000", FormulaRule(formula=[f'$M5="{val}"'], fill=PatternFill("solid", fgColor=fill)))
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = "A4:M2000"
    ws.sheet_view.showGridLines = False
    return ws


# ---------------- the same job in any tracker

STOP = {"the", "of", "and", "in", "for", "a", "an", "at", "to", "inc", "llc", "ltd", "co", "corp", "group",
        "company", "university", "&", "-", "–"}


def _words(s):
    return [w for w in re.findall(r"[a-z0-9]+", str(s or "").lower()) if w not in STOP]


def existing_index(cfg):
    """Every job already in a tracker: [{kind, sheet, row, link, emp, pos}] (Tracker and Leads sheets)."""
    out = []
    for kind, path in trackers(cfg):
        wb = load_workbook(path, read_only=True)
        for sheet in ("Tracker", "Leads"):
            if sheet not in wb.sheetnames:
                continue
            H = None
            for r, vals in enumerate(wb[sheet].iter_rows(min_row=4, values_only=True), start=4):
                if r == 4:
                    H = {h: i for i, h in enumerate(vals) if h}
                    continue
                g = lambda h: vals[H[h]] if h in H and H[h] < len(vals) else None
                emp = g("Employer")
                if not emp or str(emp).startswith("EXAMPLE"):
                    continue
                also = re.findall(r"https?://[^\s,;]+", str(g("Flags") or ""))
                out.append({"kind": kind, "sheet": sheet, "row": r, "link": norm_link(g("Link")),
                            "also": {norm_link(u) for u in also},
                            "emp": _words(emp), "pos": set(_words(g("Position"))), "decision": g("Decision")})
    return out


def find_existing(index, link, employer, position):
    """Entries in `index` that are the same job: same link, or same employer and a near-identical title."""
    link, emp, pos = norm_link(link), _words(employer), set(_words(position))
    hits = []
    for e in index:
        if link and (e["link"] == link or link in e.get("also", ())):
            hits.append(e)
            continue
        same_emp = emp and e["emp"] and (emp == e["emp"] or " ".join(emp) in " ".join(e["emp"])
                                         or " ".join(e["emp"]) in " ".join(emp))
        if same_emp and pos and e["pos"] and len(pos & e["pos"]) / len(pos | e["pos"]) >= 0.6:
            hits.append(e)
    return hits


def where(e):
    name = {"academic": "academic", "industry": "industry"}[e["kind"]]
    return f"{name} {'tracker' if e['sheet'] == 'Tracker' else 'Leads'} row {e['row']}"


def last_row(ws):
    r = ws.max_row
    while r >= 5 and all(ws.cell(row=r, column=c).value in (None, "") for c in range(1, len(COLS) + 1)):
        r -= 1
    return r


def snapshot(ws, start, end):
    """Rows start..end as lists of (value, style, hyperlink target), one per Leads column."""
    rows = []
    for r in range(start, end + 1):
        row = []
        for c in range(1, len(COLS) + 1):
            cell = ws.cell(row=r, column=c)
            row.append((cell.value, copy(cell._style), cell.hyperlink.target if cell.hyperlink else None))
        rows.append(row)
    return rows


def found_key(row):
    v = row[0][0]
    if isinstance(v, dt.datetime):
        v = v.date()
    return v.toordinal() if isinstance(v, dt.date) else 0


def write_sorted(ws, rows):
    """Write rows back from row 5, newest Found date first (stable within a date)."""
    rows = sorted(rows, key=found_key, reverse=True)
    for r, row in enumerate(rows, start=5):
        for c, (v, style, link) in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c)
            cell.value = v           # ws.cell(value=None) would leave the old value in place
            cell._style = copy(style)
            cell.hyperlink = link


def sort_leads(cfg, path):
    wb = load_workbook(path)
    if "Leads" not in wb.sheetnames:
        print("No Leads sheet yet; nothing to sort.")
        return
    ws = wb["Leads"]
    write_sorted(ws, snapshot(ws, 5, last_row(ws)))
    restore_dropdowns(wb, cfg["letter_writers"])
    wb.save(path)
    print(f"sorted {last_row(ws) - 4} lead(s), newest first, in {path}")


def add(cfg, path, items, kind="academic"):
    index = existing_index(cfg)
    wb = load_workbook(path)
    ws = ensure_sheet(wb, cfg.get("accent_color", "0021A5"), kind)
    keep, skipped = [], []
    for x in items:
        hits = find_existing(index, x.get("url"), x.get("employer"), x.get("position"))
        mine = [e for e in hits if e["kind"] == kind]
        tracked = [e for e in hits if e["kind"] != kind and e["sheet"] == "Tracker"]
        # same job already in this tracker (same link, or same employer and title), or in the other Tracker
        if mine or tracked:
            skipped.append(f"{x.get('employer')} | {x.get('position')}: already in "
                           + ", ".join(where(e) for e in (tracked or mine)))
            continue
        other = [e for e in hits if e["kind"] != kind]
        if other:
            fl = x.get("flags") or ""
            fl = "; ".join(fl) if isinstance(fl, list) else fl
            x["flags"] = "possibly the same job as " + ", ".join(where(e) for e in other) + ("; " + fl if fl else "")
        keep.append(x)
    items = keep
    first = cfg.get("scan", {}).get("priority_types", [])
    items.sort(key=lambda x: (FIT_ORDER.get(x.get("fit"), 3), x.get("type") not in first, x.get("deadline") or "9999"))
    end = last_row(ws)
    old = snapshot(ws, 5, end)
    r = end + 1          # new rows are written below, then everything is re-sorted newest first
    today = dt.date.today()
    for x in items:
        dl = None
        flags = x.get("flags") or ""
        if isinstance(flags, list):
            flags = "; ".join(str(f) for f in flags)
        if x.get("deadline"):
            m = re.search(r"\d{4}-\d{2}-\d{2}", str(x["deadline"]))
            dl = dt.date.fromisoformat(m.group(0)) if m else None
            if not m or m.group(0) != str(x["deadline"]).strip():   # e.g. "Review begins 2026-10-30"
                flags = (flags + "; " if flags else "") + "deadline note: " + str(x["deadline"])
        also = x.get("also_at") or []
        if isinstance(also, str):      # a single link or "a; b" string, not a list
            also = [s.strip() for s in re.split(r"[;,]\s*", also) if s.strip()]
        if also:
            flags = (flags + "; " if flags else "") + "also posted at " + ", ".join(also)
        vals = [today, x.get("fit"), x.get("employer"), x.get("position"),
                x.get("category") if kind == "industry" else x.get("track"), x.get("type"),
                x.get("location"), dl, x.get("why"), flags, x.get("source"), x.get("url"), None]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=FONT, size=10)
            c.border = BD
            c.alignment = Alignment(wrap_text=True, vertical="top",
                                    horizontal="center" if i in (1, 2, 5, 8, 11, 13) else None)
        ws.cell(row=r, column=1).number_format = "mmm d"
        ws.cell(row=r, column=8).number_format = "mmm d, yyyy"
        if x.get("url"):
            ws.cell(row=r, column=12).hyperlink = x["url"]
            ws.cell(row=r, column=12).font = Font(name=FONT, size=10, color="0563C1", underline="single")
        r += 1
    new = snapshot(ws, end + 1, r - 1)
    write_sorted(ws, new + old)
    restore_dropdowns(wb, cfg["letter_writers"])
    wb.save(path)
    print(f"added {len(items)} lead(s) to the top of {path}")
    for s_ in skipped:
        print("  skipped (already tracked):", s_)


PLATFORMS = [("econjobmarket", "EconJobMarket"), ("academicjobsonline", "AcademicJobsOnline"),
             ("interfolio", "Interfolio"), ("chronicle.com", "Chronicle Jobs"),
             ("insidehighered", "Inside Higher Ed Careers"), ("usajobs", "USAJOBS")]


def ad_texts(cfg, kind="academic"):
    """url -> ad text from the last scan, used to guess where to apply."""
    import os
    p = os.path.join(os.path.dirname(tracker_path(cfg, kind)), "industry_scan_new.json" if kind == "industry" else "scan_new.json")
    try:
        return {x["url"]: x.get("text", "") for x in json.load(open(p, encoding="utf-8"))["candidates"]}
    except (OSError, ValueError, KeyError):
        return {}


IND_PLATFORMS = [("myworkdayjobs", "Workday"), ("greenhouse", "Greenhouse"), ("lever.co", "Lever"),
                 ("ashbyhq", "Ashby"), ("icims", "iCIMS"), ("taleo", "Taleo"), ("linkedin.com", "LinkedIn Easy Apply"),
                 ("indeed.", "Indeed")]
SOURCES = {"LI": "LinkedIn", "IND": "Indeed", "NABE": "NABE", "JOE": "JOE"}


def apply_via(url, text, kind="academic"):
    if kind == "industry":
        u = (url or "").lower()
        return next((name for key, name in IND_PLATFORMS if key in u), "Employer website")
    t = (text or "").lower()
    for key, name in PLATFORMS[:3] + PLATFORMS[5:]:   # where the ad says to apply
        if key in t:
            return name
    if "econjobmarket" in (url or ""):
        return "EconJobMarket"
    return "Employer website"      # Chronicle / IHE ads send you to the employer's site


def sync_added(lead, tr):
    """Copy the Tracker's deadline back to Leads rows already moved there (matched by link),
    so both sheets show the date the user set in the Tracker."""
    H = {tr.cell(row=4, column=c).value: c for c in range(1, tr.max_column + 1) if tr.cell(row=4, column=c).value}
    dl = {}
    for r in range(5, tr.max_row + 1):
        link = tr.cell(row=r, column=H["Link"]).value
        if link:
            dl[str(link).strip()] = tr.cell(row=r, column=H["Deadline"]).value
    n = 0
    for r in range(5, lead.max_row + 1):
        link = lead.cell(row=r, column=12).value
        if lead.cell(row=r, column=13).value == "Added" and link and str(link).strip() in dl:
            new = dl[str(link).strip()]
            if new and lead.cell(row=r, column=8).value != new:
                lead.cell(row=r, column=8, value=new).number_format = "mmm d, yyyy"
                n += 1
    return n


def promote(cfg, path, kind="academic"):
    index = existing_index(cfg)
    wb = load_workbook(path)
    if "Leads" not in wb.sheetnames:
        print("No Leads sheet yet; nothing to promote.")
        return []
    lead, tr = wb["Leads"], wb["Tracker"]
    H = {tr.cell(row=4, column=c).value: c for c in range(1, tr.max_column + 1) if tr.cell(row=4, column=c).value}
    nxt = 5
    while tr.cell(row=nxt, column=H["Employer"]).value not in (None, ""):
        nxt += 1
    texts = ad_texts(cfg, kind)
    moved = []
    type_map = {"Tenure-track": "Tenure-track", "Teaching-focused": "Teaching-focused",
                "Fed / Central bank": "Fed / Central bank", "Government": "Government",
                "Think tank / Research": "Think tank / Research", "Consulting": "Consulting", "Tech": "Tech",
                "Finance": "Finance"}
    for r in range(5, lead.max_row + 1):
        if str(lead.cell(row=r, column=13).value or "").strip().lower() != "add":
            continue
        g = lambda c: lead.cell(row=r, column=c).value
        other = [e for e in find_existing(index, g(12), g(3), g(4)) if e["kind"] != kind and e["sheet"] == "Tracker"]
        if other:      # never put the same job in both trackers
            lead.cell(row=r, column=13, value="Added")
            lead.cell(row=r, column=10, value=f"already in {where(other[0])}; " + str(g(10) or ""))
            print(f"not promoted: {g(3)} | {g(4)} is already in {where(other[0])}")
            continue
        text = texts.get(g(12), "")
        if kind == "industry":
            # most industry jobs do not ask for letters; Yes only if the ad mentions recommendation letters
            letters = "Yes" if re.search(r"letters? of (recommendation|reference)|recommendation letter|reference letter|"
                                         r"推荐信", text + " " + str(g(10) or ""), re.I) else "No"
        else:
            # academic and policy ads nearly always want letters; say No only if a full ad never mentions them
            letters = "No" if len(text) > 1500 and not re.search(r"letter|referee|reference|recommend", text, re.I) else "Yes"
        # Letters Due only when the ad says by when materials should be received; a review date alone is just the Deadline
        received = re.search(r"deadline note:[^;|]*(received|consideration|letters? by|submitted by)", str(g(10) or ""), re.I)
        deadline, due = g(8), (g(8) if received else None)
        if due and letters == "Yes" and "Letters Due" not in H:
            H["Letters Due"] = ensure_date_column(tr, "Letters Due")
        row = {"Track": g(5), "Employer": g(3), "Position": g(4), "Type": type_map.get(g(6), "Other"),
               "Apply Via": apply_via(g(12), text, kind), "Link": g(12), "Deadline": deadline,
               "Letters Due": due if letters == "Yes" else None, "Status": "Not started",
               "Cover Letter": "To write", "Letters?": letters,
               "Notes": (f"From scan ({g(2)} fit): {g(9) or ''} | {g(10) or ''} | Apply Via and Letters? "
                         f"guessed from the ad; check.").strip()}
        if kind == "industry":
            row.update({"Track": None, "Category": g(5), "Type": g(6) or "Other", "Location": g(7),
                        "Source": SOURCES.get(g(11), "Company site"), "Resume": "To tailor", "Cover Letter": None,
                        "Form Answers": "To write"})
        for h, v in row.items():
            if h in H and v not in (None, ""):
                c = tr.cell(row=nxt, column=H[h], value=v)
                c.font = Font(name=FONT, size=10)
        for h in ("Deadline", "Letters Due"):
            if h in H and row[h]:
                tr.cell(row=nxt, column=H[h]).number_format = "mmm d, yyyy"
        if g(12):
            tr.cell(row=nxt, column=H["Link"]).hyperlink = g(12)
            tr.cell(row=nxt, column=H["Link"]).font = Font(name=FONT, size=10, color="0563C1", underline="single")
        lead.cell(row=r, column=13, value="Added")
        moved.append(g(3))
        nxt += 1
        while tr.cell(row=nxt, column=H["Employer"]).value not in (None, ""):   # never write over a row
            nxt += 1
    synced = sync_added(lead, tr)
    restore_dropdowns(wb, cfg["letter_writers"])
    wb.save(path)
    print(f"promoted {len(moved)}: {moved}; synced {synced} deadline(s) back to Leads")
    return moved


if __name__ == "__main__":
    cfg = load_config()
    kind = "industry" if "--industry" in sys.argv else "academic"
    sys.argv = [a for a in sys.argv if a != "--industry"]
    tracker = tracker_path(cfg, kind)
    if not os.path.exists(tracker):
        raise SystemExit(f"{tracker} does not exist yet. Build it with: python3 scripts/build_tracker.py"
                         + (" --industry" if kind == "industry" else ""))
    cmd = sys.argv[1] if len(sys.argv) >= 2 else ""
    if cmd not in ("add", "promote", "sort") or (cmd == "add" and len(sys.argv) < 3):
        print(__doc__)
        sys.exit()
    # Excel keeps showing (and may later save) its in-memory copy, so save and close it first
    was_open = close_in_excel(tracker)
    if cmd == "add":
        add(cfg, tracker, json.load(open(sys.argv[2], encoding="utf-8")), kind)
    elif cmd == "promote":
        promote(cfg, tracker, kind)
    else:
        sort_leads(cfg, tracker)
    reopen_in_excel(tracker, was_open)
