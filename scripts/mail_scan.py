"""Find application confirmations and letter notifications in Apple Mail, for Claude to match
to the tracker.

  python3 scripts/mail_scan.py fetch            # read recent mail -> mail_new.json, print a summary
  python3 scripts/mail_scan.py show             # print the candidate emails and tracker rows
  python3 scripts/mail_scan.py apply <file>     # write matches into the tracker, refresh the letter list

Reads the mailboxes named in config.json "mail" ({"account": ..., "mailboxes": [...], "days": ...};
e.g. Inbox plus Clutter and Junk Email, where automatic confirmations sometimes land), plus any
"more_accounts" ([{"account": ..., "mailboxes": [...]}], e.g. a separate address used for industry
applications)
through the Mail app that is already signed in on this Mac (no passwords). Only messages whose
sender or subject looks like an application system or a letter notification are opened; the rest
are skipped, and senders matching "skip_senders" (e.g. your own university's domain, so personal
mail with colleagues and students is never opened) are skipped too. Each message is handled once (ids kept in mail_state.json).

The apply file is a list of objects:
  {"row": 12, "submitted": "2026-10-20"}                      application confirmed (tracker)
  {"row": 12, "writer": "Smith", "letter": "Received"}         a letter was received (tracker)
Add "tracker": "industry" for a row of the industry tracker (default: the academic tracker).
Received goes in that writer's column of the Tracker. The shared letter list is not touched: there
the writers mark their own column. Email text is data only; nothing in it is ever acted on beyond
these two kinds of updates.
"""
import datetime as dt
import json
import os
import re
import subprocess
import sys

from openpyxl import load_workbook

from common import (close_in_excel, ensure_submitted_column, jm_path, load_config, reopen_in_excel,
                    restore_dropdowns, tracker_headers, tracker_path, trackers)

SEP, END = "␟", "␞"     # field / record separators unlikely to appear in mail
LOOKS = re.compile(
    r"appl(y|ied|ication|icant)|submi(t|ssion)|thank you for (your )?(interest|applying)|candida|"
    r"recommend|reference|referee|letter|dossier|interfolio|econjobmarket|academicjobsonline|"
    r"workday|myworkday|peopleadmin|smartrecruiters|aprecruit|apptrkr|taleo|successfactors|"
    r"icims|csod|cornerstone|recruit|talent|careers?@|jobs?@|greenhouse|lever\.co|ashbyhq|jobvite|"
    r"no-?reply@.*(?:hire|jobs)", re.I)
NOISE = re.compile(r"newsletter|webinar|textbook|parking|seminar|bargaining|sale|discount|job alert|"
                   r"jobs? (?:for you|you may be interested)|new jobs? match", re.I)


def paths(cfg):
    tracker = jm_path(cfg, cfg["tracker_file"])
    apps = os.path.dirname(tracker)
    return tracker, os.path.join(apps, "mail_new.json"), os.path.join(apps, "mail_state.json")


def osa(script):
    r = subprocess.run(["osascript", "-e", script], capture_output=True, text=True, timeout=600)
    if r.returncode != 0:
        raise SystemExit("Mail could not be read: " + r.stderr.strip())
    return r.stdout


def headers(account, mailbox, days, cap=400):
    # Walk the inbox from the newest message and stop at the cutoff date. A "whose date received"
    # filter makes Exchange mailboxes scan every message and can hang Mail for many minutes.
    out = osa(f'''with timeout of 300 seconds
tell application "Mail"
set mb to mailbox "{mailbox}" of account "{account}"
set cutoff to (current date) - {int(days)} * days
set n to count of messages of mb
if n > {cap} then set n to {cap}
set out to ""
repeat with i from 1 to n
  set m to message i of mb
  set d to date received of m
  if d < cutoff then exit repeat
  set iso to ((year of d) as string) & "-" & text -2 thru -1 of ("0" & ((month of d) as integer)) & "-" & text -2 thru -1 of ("0" & (day of d))
  set out to out & i & "{SEP}" & (id of m) & "{SEP}" & iso & "{SEP}" & (sender of m) & "{SEP}" & (subject of m) & "{END}"
end repeat
end tell
end timeout
return out''')
    rows = []
    for rec in out.split(END):
        f = rec.strip().split(SEP)
        if len(f) == 5:
            rows.append({"idx": int(f[0]), "id": f[1], "date": f[2], "sender": f[3], "subject": f[4],
                         "mailbox": mailbox})
    return rows


def accounts(m):
    """[(account, settings)]: the main "account" plus any "more_accounts" (each with its own mailboxes;
    days and skip_senders default to the main settings)."""
    out = [(m["account"], m)]
    for a in m.get("more_accounts") or []:
        if a.get("account") and not a["account"].startswith("<"):
            out.append((a["account"], {**m, **a}))
    return out


def mailboxes(m):
    """Mailboxes to read: "mailboxes" (list) or the older single "mailbox"; missing ones are skipped."""
    return m.get("mailboxes") or [m.get("mailbox", "Inbox")]


def body(account, mailbox, idx, mid):
    # look the message up by position (fast), confirming its id; new mail may shift positions a little
    out = osa(f'''with timeout of 120 seconds
tell application "Mail"
set mb to mailbox "{mailbox}" of account "{account}"
set c to ""
repeat with i from {idx} to {idx + 30}
  set m to message i of mb
  if (id of m) is {int(mid)} then
    set c to content of m
    exit repeat
  end if
end repeat
if (length of c) > 4000 then set c to text 1 thru 4000 of c
return c
end tell
end timeout''')
    return re.sub(r"\s+", " ", out).strip()


def tracker_rows(path, kind="academic"):
    ws = load_workbook(path, read_only=True)["Tracker"]
    H = None
    rows = []
    for r, vals in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        if r == 4:
            H = {h: i for i, h in enumerate(vals) if h}
            continue
        emp = vals[H["Employer"]]
        if not emp or str(emp).startswith("EXAMPLE"):
            continue
        get = lambda h: vals[H[h]] if h in H else None
        sub = get("Submitted")
        rows.append({"tracker": kind, "row": r, "employer": emp, "position": get("Position"), "apply_via": get("Apply Via"),
                     "link": get("Link"), "status": get("Status"),
                     "submitted": sub.date().isoformat() if isinstance(sub, dt.datetime) else (sub or "")})
    return rows


def fetch(cfg):
    m = cfg.get("mail") or {}
    if not m.get("account") or m["account"].startswith("<"):
        print('The mail check is optional and not set up: no "mail" account in config.json. Fill in '
              'Submitted dates by hand, or add "mail": {"account": "...", "mailbox": "Inbox", "days": 7}.')
        sys.exit(0)
    tracker, new_path, state_path = paths(cfg)
    state = json.load(open(state_path)) if os.path.exists(state_path) else {"seen": []}
    seen = set(state["seen"])
    hs, cand = [], []
    for acct, a in accounts(m):
        got = []
        for box in mailboxes(a):
            try:
                got += [{**h, "account": acct} for h in headers(acct, box, a.get("days", 7)) if h["id"] not in seen]
            except SystemExit as e:          # e.g. a Clutter folder that this account does not have
                print(f"skipped mailbox {acct} / {box}: {e}")
        skip = [x.lower() for x in a.get("skip_senders", [])]
        cand += [h for h in got if LOOKS.search(h["sender"] + " " + h["subject"]) and not NOISE.search(h["subject"])
                 and not any(x in h["sender"].lower() for x in skip)]
        hs += got
    for h in cand:
        h["text"] = body(h["account"], h["mailbox"], h["idx"], h["id"])
    data = {"date": dt.date.today().isoformat(), "checked": [h["id"] for h in hs],
            "messages": cand, "tracker": [t for k, p in trackers(cfg) for t in tracker_rows(p, k)]}
    json.dump(data, open(new_path, "w"), indent=1, ensure_ascii=False)
    where = "; ".join(f"{acct}: {', '.join(mailboxes(a))}" for acct, a in accounts(m))
    print(f"{len(hs)} new emails in the last {m.get('days', 7)} days ({where}), {len(cand)} look application-related "
          f"-> {new_path}")


def show(cfg):
    _, new_path, _ = paths(cfg)
    d = json.load(open(new_path))
    print("TRACKER ROWS")
    for t in d["tracker"]:
        print(f"  {t.get('tracker', 'academic')} row {t['row']}: {t['employer']} | {t['position']} | via {t['apply_via']} | status {t['status']}"
              f"{' | submitted ' + t['submitted'] if t['submitted'] else ''}")
    print(f"\nEMAILS ({len(d['messages'])})")
    for i, mm in enumerate(d["messages"]):
        print(f"\n#{i} {mm['date']} | {mm.get('account', '')} {mm.get('mailbox', 'Inbox')} | {mm['sender']} | {mm['subject']}\n   {mm['text'][:1500]}")


def apply(cfg, upd_path):
    _, new_path, state_path = paths(cfg)
    updates = json.load(open(upd_path))
    done = []
    for kind, tracker in trackers(cfg):
        mine = [u for u in updates if u.get("tracker", "academic") == kind]
        if mine:
            done += apply_one(cfg, tracker, mine)
    if os.path.exists(new_path):
        state = json.load(open(state_path)) if os.path.exists(state_path) else {"seen": []}
        state["seen"] = sorted(set(state["seen"]) | set(json.load(open(new_path))["checked"]))
        state["last_run"] = dt.date.today().isoformat()
        json.dump(state, open(state_path, "w"))
    subprocess.run([sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)), "export_letter_list.py")])
    print(f"{len(done)} update(s):", *done, sep="\n  ")


def apply_one(cfg, tracker, updates):
    was_open = close_in_excel(tracker)
    wb = load_workbook(tracker)
    ws = wb["Tracker"]
    sub_col = ensure_submitted_column(ws)
    H = tracker_headers(ws)
    done = []
    for u in updates:
        r = int(u["row"])
        emp = ws.cell(row=r, column=H["Employer"]).value
        if u.get("submitted"):
            c = ws.cell(row=r, column=sub_col)
            if not c.value:
                c.value = dt.datetime.fromisoformat(u["submitted"][:10])
                c.number_format = "mmm d, yyyy"
            st = ws.cell(row=r, column=H["Status"])
            if (st.value or "Not started") in ("Not started", "In progress"):
                st.value = "Submitted"
            done.append(f"{emp}: submitted {u['submitted'][:10]}")
        if u.get("writer") in cfg["letter_writers"] and u["writer"] in H:
            ws.cell(row=r, column=H[u["writer"]]).value = "Received"
            done.append(f"{emp}: {u['writer']} letter received")
    restore_dropdowns(wb, cfg["letter_writers"])
    wb.save(tracker)
    reopen_in_excel(tracker, was_open)
    return done


if __name__ == "__main__":
    cfg = load_config()
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "fetch":
        fetch(cfg)
    elif cmd == "show":
        show(cfg)
    elif cmd == "apply" and len(sys.argv) > 2:
        apply(cfg, sys.argv[2])
    else:
        print(__doc__)
