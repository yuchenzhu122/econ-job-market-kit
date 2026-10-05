"""Pull new academic/policy postings from JOE, EconJobMarket, Chronicle and Inside Higher Ed and pre-filter them.

  python3 scripts/scan_postings.py            # writes <job_market_dir>/08_Applications/scan_new.json
  python3 scripts/scan_postings.py --all      # ignore the 'already seen' list (for testing)
  python3 scripts/scan_postings.py show [start] [count]   # compact list of candidates for triage
  python3 scripts/scan_postings.py text <n> [<n> ...]     # full text of candidates by number

Sources (all public, machine-readable):
  JOE: the AEA's XML export of current listings
  EJM: EconJobMarket's public JSON feed of ads
  CHE / IHE: Chronicle of Higher Education Jobs and Inside Higher Ed Careers (RSS search
  feeds, newest first; strong on teaching-focused and regional colleges). HigherEdJobs
  blocks automated access, so it is not scanned; most of its economics faculty ads are
  cross-posted to one of these.

Rule-based pre-filter (settings in config.json "scan"): section / position type, rank
(drops senior-only and postdoc/visiting), field (JEL prefixes, categories, keywords),
deadline not passed, not already seen, not already in the tracker. What survives is
written to scan_new.json for Claude to read in full and judge fit (see skills/econ-job-scan).
"""
import datetime as dt
import html
import json
import os
import re
import sys
import urllib.request

from openpyxl import load_workbook

from common import jm_path, load_config

UA = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/129 Safari/537.36 (econ-job-market-kit scanner)"}
JOE_LIST = "https://www.aeaweb.org/joe/listings"
EJM_JSON = "https://backend.econjobmarket.org/data/zz_public/json/Ads"
RSS_SITES = {"CHE": "https://jobs.chronicle.com/jobsrss/?PositionType=12",           # Economics category
             "IHE": "https://careers.insidehighered.com/jobsrss/?FacultyJobs=57"}  # Economics faculty
RSS_DAYS = 120      # ignore RSS ads posted longer ago than this (boards keep old pooled ads)
RSS_PAGES = 20      # 20-25 ads per page

JUNIOR = re.compile(r"assistant prof|lecturer|instructor|teaching|tenure[- ]track|open rank|economist|"
                    r"research (?:associate|fellow|scientist)|policy|analyst|all ranks|any rank", re.I)
EXCLUDE = re.compile(r"post-?doc|visiting|pre-?doc|research assistant|phd (?:student|position|fellowship)|"
                     r"doctoral (?:student|fellowship)|adjunct|part-time|dean|chair\b|director|head of", re.I)
VISA = re.compile(r"[^.]*(?:sponsor|citizen|permanent resident|green card|work authori[sz]ation|"
                  r"authori[sz]ed to work|right to work|visa|security clearance)[^.]*\.", re.I)
OTHER_FIELD = re.compile(r"financ|marketing|accounting|management|philosoph|real estate|business law|"
                         r"information systems|entrepreneur|supply chain|operations", re.I)


def fetch(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=60) as r:
        return r.read().decode("utf-8", errors="ignore")


def clean(s):
    s = html.unescape(re.sub(r"<[^>]+>", " ", s or ""))
    return re.sub(r"\s+", " ", s).strip()


def joe_postings():
    page = fetch(JOE_LIST)
    m = re.search(r'href="(/joe/resultset_output\.php\?mode=full_xml[^"]*)"', page)
    if not m:
        raise RuntimeError("JOE export link not found")
    xml = fetch("https://www.aeaweb.org" + html.unescape(m.group(1)))
    # positions sit inside <year joe_year_ID="2026"><issue joe_issue_ID="2">; the listing URL
    # is JOE_ID=<year>-<2-digit issue>_<jp_id>
    marks = [(m.start(), "year" if m.group(1) else "issue", m.group(2) or m.group(4)) for m in
             re.finditer(r'<(year) joe_year_ID="(\d+)"|<(issue) joe_issue_ID="(\d+)"', xml)]
    out = []
    for p in re.finditer(r"<position jp_id=\"(\d+)\">(.*?)</position>", xml, re.S):
        jid, body = p.group(1), p.group(2)
        year = next((v for pos, k, v in reversed(marks) if k == "year" and pos < p.start()), "")
        issue = next((v for pos, k, v in reversed(marks) if k == "issue" and pos < p.start()), "")
        g = lambda t: clean((re.search(rf"<{t}>(.*?)</{t}>", body, re.S) or [None, ""])[1])
        countries = re.findall(r"<country>(.*?)</country>", body)
        cities = re.findall(r"<city>(.*?)</city>", body)
        out.append({
            "source": "JOE", "id": "JOE-" + jid,
            "url": (f"https://www.aeaweb.org/joe/listing.php?JOE_ID={year}-{int(issue):02d}_{jid}"
                    if year and issue else "https://www.aeaweb.org/joe/listings"),
            "section": g("jp_section"), "title": g("jp_title"),
            "employer": g("jp_institution"),
            "department": " / ".join(x for x in (g("jp_division"), g("jp_department")) if x),
            "location": ", ".join(x for x in (clean(cities[0]) if cities else "",
                                              clean(countries[0]) if countries else "") if x),
            "deadline": g("jp_application_deadline")[:10],
            "fields": [c for c in re.findall(r"<jc_code>(.*?)</jc_code>", body)],
            "field_names": [clean(c) for c in re.findall(r"<jc_description>(.*?)</jc_description>", body)],
            "text": g("jp_full_text"),
        })
    return out


def ejm_postings():
    out = []
    for a in json.loads(fetch(EJM_JSON)):
        loc = (a.get("locations") or [{}])[0] or {}
        out.append({
            "source": "EJM", "id": "EJM-" + a["url"].rstrip("/").split("/")[-1], "url": a["url"],
            "section": "; ".join(t["name"] for t in a.get("position_types") or []),
            "title": clean(a.get("adtitle")), "employer": clean(a.get("name")),
            "department": clean(a.get("department")),
            "location": ", ".join(x for x in (loc.get("city"), loc.get("country_code")) if x),
            "deadline": (a.get("deadline_date") or a.get("enddate") or "")[:10],
            "end": (a.get("enddate") or "")[:10],
            "fields": [], "field_names": [c["name"] for c in a.get("categories") or []],
            "text": clean(a.get("adtext")),
        })
    return out


def rss_postings(src):
    """Ads in the Economics category of a Madgex job board's RSS feed. Full text is fetched
    later (rss_details) only for ads that pass the filter, to keep requests few."""
    import email.utils
    out, seen = [], set()
    cutoff = dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=RSS_DAYS)
    for page in range(1, RSS_PAGES + 1):
        items = re.findall(r"<item>(.*?)</item>", fetch(f"{RSS_SITES[src]}&page={page}"), re.S)
        if not items:
            break
        for it in items:
            g = lambda t: html.unescape((re.search(rf"<{t}>(.*?)</{t}>", it, re.S) or [None, ""])[1]).strip()
            link = g("link").split("?")[0]
            m = re.search(r"/job/(\d+)/", link)
            if not m or m.group(1) in seen:
                continue
            seen.add(m.group(1))
            try:
                if email.utils.parsedate_to_datetime(g("pubDate")) < cutoff:
                    continue
            except (TypeError, ValueError):
                pass
            emp, _, title = clean(g("title")).partition(": ")
            if not title:
                emp, title = "", emp
            out.append({"source": src, "id": f"{src}-{m.group(1)}", "url": link, "section": "Academic",
                        "title": title, "employer": emp, "department": "", "location": "", "deadline": "",
                        "fields": [], "field_names": [], "text": clean(g("description"))})
    return out


def rss_details(x):
    """Fill location and full text from the job page (Madgex layout)."""
    page = fetch(x["url"])
    dd = lambda k: clean((re.search(rf'<dt class="mds-list__key">{k}</dt>\s*<dd[^>]*>(.*?)</dd>', page, re.S)
                          or [None, ""])[1])
    x["location"] = dd("Location").replace(", United States", "")
    x["employer"] = x["employer"] or dd("Employer")
    body = re.search(r'<section class="mds-tabs__panel" id="job-description">(.*?)</section>', page, re.S)
    if body:
        x["text"] = clean(re.sub(r"<(script|style).*?</\1>", " ", body.group(1), flags=re.S))
    m = re.search(r"(?:deadline|review of applications|applications will be reviewed|priority consideration)"
                  r"[^.]{0,80}?((?:January|February|March|April|May|June|July|August|September|October|November|"
                  r"December) \d{1,2},? 20\d\d)", x["text"], re.I)
    if m:
        try:
            x["deadline"] = dt.datetime.strptime(m.group(1).replace(",", ""), "%B %d %Y").date().isoformat()
        except ValueError:
            pass


def keep(x, sc, today):
    why = []
    if x["source"] == "JOE":
        if x["section"] not in sc["joe_sections"]:
            return False, "section"
    elif x["source"] in RSS_SITES:
        pass        # Economics category of academic boards; fields judged later from full text
    else:
        types = [t.strip() for t in x["section"].split(";")]
        if not any(t in sc["ejm_position_types"] for t in types):
            return False, "type"
        if x.get("end") and x["end"] < today:
            return False, "expired"
    head = x["title"] + " " + x["section"]
    if EXCLUDE.search(x["title"]) and not re.search(r"assistant prof|lecturer", x["title"], re.I):
        return False, "rank"
    if re.search(r"adjunct|part\s*-?\s*time|\bPT\b", x["title"], re.I):
        return False, "rank"
    if (x["source"] in RSS_SITES and OTHER_FIELD.search(x["title"])
            and not re.search(r"econ", x["title"], re.I)):
        return False, "field"
    generic = x["source"] in RSS_SITES and re.search(r"faculty|professor", x["title"], re.I)
    if not (JUNIOR.search(head) or generic) and "Assistant Professor" not in x["section"] and "Lecturer" not in x["section"]:
        return False, "rank"
    if x["deadline"] and x["deadline"] < today and x["source"] == "JOE":
        return False, "deadline"
    if x["source"] in RSS_SITES:
        return True, ""
    text = (x["text"] + " " + " ".join(x["field_names"]) + " " + x["title"]).lower()
    field_ok = (any(f.startswith(tuple(sc["field_jel_prefixes"])) for f in x["fields"])
                or any(c in sc["ejm_categories"] for c in x["field_names"])
                or any(k in text for k in sc["field_keywords"]))
    if not field_ok:
        return False, "field"
    return True, ""


def main():
    cfg = load_config()
    sc = cfg["scan"]
    apps = jm_path(cfg, os.path.dirname(cfg["tracker_file"]))
    state_path = os.path.join(apps, "scan_state.json")
    state = json.load(open(state_path)) if os.path.exists(state_path) else {"seen": []}
    seen = set() if "--all" in sys.argv else set(state["seen"])

    tracked = set()
    tracker = jm_path(cfg, cfg["tracker_file"])
    if os.path.exists(tracker):
        wb = load_workbook(tracker, read_only=True)
        for name in ("Tracker", "Leads"):
            if name in wb.sheetnames:
                for row in wb[name].iter_rows(values_only=True):
                    for v in row:
                        if isinstance(v, str) and v.startswith("http"):
                            tracked.add(v.strip())

    today = dt.date.today().isoformat()
    posts, errors = [], []
    for name, fn in (("JOE", joe_postings), ("EJM", ejm_postings),
                     ("CHE", lambda: rss_postings("CHE")), ("IHE", lambda: rss_postings("IHE"))):
        try:
            posts += fn()
        except Exception as e:  # keep going if one source is down
            errors.append(f"{name}: {e}")

    stats, new = {}, []
    for x in posts:
        if x["id"] in seen or x["url"] in tracked:
            stats["seen"] = stats.get("seen", 0) + 1
            continue
        ok, reason = keep(x, sc, today)
        if ok and x["source"] in RSS_SITES:
            try:
                rss_details(x)
            except Exception as e:
                errors.append(f"{x['source']} page {x['url']}: {e}")
            if x["deadline"] and x["deadline"] < today:
                ok, reason = False, "deadline"
        if ok:
            # work-authorization sentences, kept even when the ad text is cut below
            x["visa_text"] = " ".join(dict.fromkeys(m.strip() for m in VISA.findall(x["text"])))[:800]
            x["text"] = x["text"][:5000]
            x.pop("end", None)
            new.append(x)
        else:
            stats[reason] = stats.get(reason, 0) + 1

    # the same job is often posted on several boards: keep one (JOE > EJM > CHE > IHE), remember the other link
    def key(x):
        emp = re.sub(r"\(.*?\)", "", x["employer"].lower())
        emp = re.sub(r"[^a-z]", "", emp.replace("university", "").replace("the", ""))[:20]
        ttl = " ".join([w for w in re.findall(r"[a-z]+", x["title"].lower()) if w not in ("of", "in", "and", "the")][:3])
        return emp + "|" + ttl
    merged = {}
    order = {"JOE": 0, "EJM": 1, "CHE": 2, "IHE": 3}
    for x in sorted(new, key=lambda x: order.get(x["source"], 9)):
        k = key(x)
        if k in merged:
            merged[k].setdefault("also_at", []).append(x["url"])
            stats["duplicate"] = stats.get("duplicate", 0) + 1
        else:
            merged[k] = x
    new = sorted(merged.values(), key=lambda x: x["deadline"] or "9999")
    for x in new:
        x["summary"] = x["text"][:1200]

    out_path = os.path.join(apps, "scan_new.json")
    json.dump({"date": today, "errors": errors, "filtered_out": stats, "candidates": new},
              open(out_path, "w"), indent=1, ensure_ascii=False)
    # remember everything we looked at, so the next run only shows new postings
    state["seen"] = sorted(set(state["seen"]) | {x["id"] for x in posts})
    state["last_run"] = today
    json.dump(state, open(state_path, "w"))
    print(f"{len(posts)} postings fetched, {len(new)} new candidates -> {out_path}")
    print("filtered out:", stats, "| errors:", errors or "none")


def view():
    """Read scan_new.json in pieces, so Claude never needs ad-hoc scripts (each needs approval)."""
    cfg = load_config()
    path = os.path.join(jm_path(cfg, os.path.dirname(cfg["tracker_file"])), "scan_new.json")
    data = json.load(open(path, encoding="utf-8"))
    c = data["candidates"]
    if sys.argv[1] == "show":
        start = int(sys.argv[2]) if len(sys.argv) > 2 else 0
        count = int(sys.argv[3]) if len(sys.argv) > 3 else 40
        if start == 0:
            print(f"scan date {data['date']} | {len(c)} candidates | errors: {data['errors'] or 'none'} "
                  f"| filtered out: {data['filtered_out']}\n")
        for n in range(start, min(start + count, len(c))):
            x = c[n]
            print(f"#{n} [{x['source']}] {x['employer']} | {x['title']} | {x['section']} | {x['location']} | "
                  f"deadline {x['deadline'] or '?'} | fields: {', '.join(x['field_names'][:6]) or '-'}")
            if x.get("visa_text"):
                print(f"   VISA: {x['visa_text'][:300]}")
            print(f"   {x['summary'][:500]}\n")
        if start + count < len(c):
            print(f"... next: show {start + count} {count}")
    else:
        for n in sys.argv[2:]:
            x = c[int(n)]
            print(f"#{n} [{x['source']}] {x['employer']} | {x['title']} | {x['url']}")
            print(f"also at: {x.get('also_at') or '-'} | deadline {x['deadline'] or '?'} | {x['location']}")
            print(f"VISA: {x.get('visa_text') or '(nothing in ad)'}")
            print(x["text"] + "\n" + "-" * 60)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] in ("show", "text"):
        view()
    else:
        main()
