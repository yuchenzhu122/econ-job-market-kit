"""Find industry postings for economics PhDs: job-alert emails (LinkedIn, Indeed) and company career pages.

  python3 scripts/industry_scan.py              # writes <industry tracker folder>/industry_scan_new.json
  python3 scripts/industry_scan.py --all        # ignore the 'already seen' list (for testing)
  python3 scripts/industry_scan.py --no-mail    # skip the job-alert emails
  python3 scripts/industry_scan.py --companies  # also read company career pages when scan_companies is false
  python3 scripts/industry_scan.py --board-file nabe_jobs.json   # add postings read in the browser
                                                (NABE EconJobs: see scripts/nabe_jobs.js)
  python3 scripts/industry_scan.py show [start] [count]   # compact list of candidates for triage
  python3 scripts/industry_scan.py text <n> [<n> ...]     # full text of candidates by number

Settings: config.json "industry":
  categories      {"Econ / Consulting": ["Economist", ...], ...}: a title must contain one of these
                  keywords; the first category that matches labels the posting
  exclude_titles  words that drop a posting (Senior, Staff, Director, Intern, ...)
  alert_mail      {"account", "mailboxes", "days", "boards"}: job-alert emails in the Mail app (no
                  passwords). Each board = {"name", "code", "senders", "link"}: only messages whose sender
                  contains one of `senders` are opened, and only links matching `link` are kept.
                  Default boards: LinkedIn, Indeed.
  scan_companies  false = skip the career pages below (alert emails are the main source)
  companies       [{"name", "ats": greenhouse | lever | ashby | workday | amazon | google | eightfold,
                    "board" or "url" (eightfold: the site, e.g. https://explore.jobs.netflix.net, plus
                    "domain"), "type", "searches" (search terms for workday / amazon / google / eightfold)}]

LinkedIn and Indeed do not allow scraping, so their postings come only from the alert emails you set
up on those sites. NABE's EconJobs board blocks scripts but is public: Claude reads it in the built-in
browser with scripts/nabe_jobs.js and passes the result with --board-file. Postings already seen, or already in any tracker (academic or industry,
Tracker or Leads sheet), are skipped. What survives goes to industry_scan_new.json for Claude to
judge (skills/econ-industry-scan).
"""
import datetime as dt
import email
import email.policy
import html
import json
import os
import re
import sys
import urllib.parse

from openpyxl import load_workbook

from common import load_config, norm_link, tracker_path, trackers
from scan_postings import VISA, clean, fetch, post_json

DEFAULT_EXCLUDE = ["senior", "sr.", "sr ", "staff", "principal", "lead", "manager", "director", "head of", "vp",
                   "vice president", "intern", "internship", "co-op", "part-time", "contract", "postdoc"]


# ---------------- company career pages

def greenhouse(c):
    d = json.loads(fetch(f"https://boards-api.greenhouse.io/v1/boards/{c['board']}/jobs?content=true"))
    return [{"id": f"GH-{c['board']}-{j['id']}", "url": j["absolute_url"], "title": j["title"],
             "location": (j.get("location") or {}).get("name", ""), "posted": (j.get("updated_at") or "")[:10],
             "text": clean(html.unescape(j.get("content") or ""))} for j in d.get("jobs", [])]


def lever(c):
    d = json.loads(fetch(f"https://api.lever.co/v0/postings/{c['board']}?mode=json"))
    return [{"id": f"LV-{c['board']}-{j['id']}", "url": j["hostedUrl"], "title": j["text"],
             "location": (j.get("categories") or {}).get("location", ""),
             "posted": dt.date.fromtimestamp(j["createdAt"] / 1000).isoformat() if j.get("createdAt") else "",
             "text": clean((j.get("descriptionPlain") or "") + " " + " ".join(
                 l.get("text", "") + " " + clean(l.get("content", "")) for l in j.get("lists") or []))} for j in d]


def ashby(c):
    d = json.loads(fetch(f"https://api.ashbyhq.com/posting-api/job-board/{c['board']}"))
    return [{"id": f"AB-{c['board']}-{j['id']}", "url": j.get("jobUrl") or j.get("applyUrl"), "title": j["title"],
             "location": j.get("location", ""), "posted": (j.get("publishedAt") or "")[:10],
             "text": clean(j.get("descriptionPlain") or j.get("descriptionHtml") or "")} for j in d.get("jobs", [])]


def workday(c, keep_title):
    """Workday's public job API, searched with each term in c["searches"]; details only for kept titles."""
    base = c["url"].rstrip("/")
    site = re.match(r"https://([^/]+)/wday/cxs/[^/]+/([^/]+)", base)
    out, seen = [], set()
    for term in c.get("searches") or ["economist", "economics", "causal", "pricing", "forecasting"]:
        d = post_json(base + "/jobs", {"appliedFacets": {}, "limit": 20, "offset": 0, "searchText": term})
        for j in d.get("jobPostings", []):
            if j["externalPath"] in seen or not keep_title(j.get("title", "")):
                continue
            seen.add(j["externalPath"])
            info = json.loads(fetch(base + j["externalPath"])).get("jobPostingInfo", {})
            out.append({"id": f"WD-{c['name']}-{info.get('jobReqId') or j['externalPath'].split('_')[-1]}",
                        "url": f"https://{site.group(1)}/en-US/{site.group(2)}{j['externalPath']}",
                        "title": j.get("title", ""), "location": j.get("locationsText", ""),
                        "posted": (info.get("startDate") or "")[:10], "text": clean(info.get("jobDescription"))})
    return out


def amazon(c):
    out, seen = [], set()
    for term in c.get("searches") or ["economist"]:
        q = urllib.parse.urlencode({"base_query": term, "result_limit": 100, "sort": "recent"})
        for j in json.loads(fetch("https://www.amazon.jobs/en/search.json?" + q)).get("jobs", []):
            if j["id_icims"] in seen:
                continue
            seen.add(j["id_icims"])
            out.append({"id": f"AMZN-{j['id_icims']}", "url": "https://www.amazon.jobs" + j["job_path"],
                        "title": j["title"], "location": j.get("normalized_location") or j.get("location", ""),
                        "posted": j.get("posted_date", ""),
                        "text": clean(" ".join(j.get(k) or "" for k in ("description", "basic_qualifications",
                                                                         "preferred_qualifications")))})
    return out


def google(c, keep_title):
    """Google's careers search page (first page only per term, as its robots.txt asks); the job
    data sits in the page's AF_initDataCallback 'ds:1' block."""
    out, seen = [], set()
    for term in c.get("searches") or ["economist"]:
        page = fetch("https://www.google.com/about/careers/applications/jobs/results?" + urllib.parse.urlencode({"q": term}))
        m = re.search(r"AF_initDataCallback\(\{key: 'ds:1'.*?data:(.*?), sideChannel", page, re.S)
        for j in (json.loads(m.group(1))[0] or []) if m else []:
            if j[0] in seen or not keep_title(j[1]):
                continue
            seen.add(j[0])
            txt = " ".join(clean((x or [None, ""])[1]) for x in (j[10], j[3], j[4], j[19]) if isinstance(x, list))
            out.append({"id": f"GOOG-{j[0]}", "title": j[1],
                        "url": f"https://www.google.com/about/careers/applications/jobs/results/{j[0]}-"
                               + re.sub(r"[^a-z0-9]+", "-", j[1].lower()).strip("-"),
                        "location": "; ".join(l[0] for l in j[9] or []), "posted": "", "text": txt})
    return out


def eightfold(c, keep_title):
    """Eightfold career sites (Netflix): /api/apply/v2, which their robots.txt allows."""
    base, out, seen = c["url"].rstrip("/"), [], set()
    for term in c.get("searches") or ["economist", "data scientist", "research scientist"]:
        q = urllib.parse.urlencode({"domain": c["domain"], "query": term, "num": 50})
        for p in json.loads(fetch(f"{base}/api/apply/v2/jobs?{q}")).get("positions", []):
            if p["id"] in seen or not keep_title(p["name"]):
                continue
            seen.add(p["id"])
            d = json.loads(fetch(f"{base}/api/apply/v2/jobs/{p['id']}?domain={c['domain']}"))
            out.append({"id": f"EF-{c['name']}-{p['id']}", "title": p["name"],
                        "url": p.get("canonicalPositionUrl") or f"{base}/careers/job/{p['id']}",
                        "location": p.get("location", ""),
                        "posted": dt.date.fromtimestamp(p["t_create"]).isoformat() if p.get("t_create") else "",
                        "text": clean(d.get("job_description"))})
    return out


# ---------------- job-alert emails

BOARDS = [   # default job-alert sources; config industry.alert_mail.boards replaces this list
    {"name": "LinkedIn", "code": "LI", "senders": ["jobalerts-noreply@linkedin.com", "jobs-noreply@linkedin.com"],
     "link": r"linkedin\.com/(?:comm/)?jobs/view/"},
    {"name": "Indeed", "code": "IND", "senders": ["indeed.com"], "link": r"indeed\.[a-z.]+/.*[?&](?:jk|vjk)="},
    # NABE EconJobs needs no alert: its board is read directly in the browser (--board-file)
]


def boards(m):
    return m.get("boards") or BOARDS


def alert_messages(m, seen):
    """Raw MIME source of recent messages from the job boards' alert senders, via the Mail app."""
    from mail_scan import headers, mailboxes, osa
    senders = [(s.lower(), b) for b in boards(m) for s in b["senders"]]
    out = []
    for box in mailboxes(m):
        try:
            hs = headers(m["account"], box, m.get("days", 7))
        except SystemExit as e:
            print(f"skipped mailbox {box}: {e}")
            continue
        for h in hs:
            board = next((b for s, b in senders if s in h["sender"].lower()), None)
            if h["id"] in seen or not board:
                continue
            src = osa(f'''with timeout of 120 seconds
tell application "Mail"
set mb to mailbox "{box}" of account "{m['account']}"
repeat with i from {h['idx']} to {h['idx'] + 30}
  set x to message i of mb
  if (id of x) is {int(h['id'])} then return source of x
end repeat
end tell
end timeout''')
            out.append({**h, "source": src, "board": board})
    return out


def unwrap(u, pattern):
    """The job link inside a tracking redirect (?url=https%3A%2F%2F...), or u itself."""
    for _ in range(3):
        if re.search(pattern, u):
            return u
        inner = [v for vals in urllib.parse.parse_qs(urllib.parse.urlparse(u).query).values() for v in vals
                 if v.startswith("http")]
        if not inner:
            return u
        u = inner[0]
    return u


AGE = re.compile(r"^(just posted|today|active \d+ days? ago|\d+\+? days? ago|posted \d+\+? days? ago)$", re.I)


def indeed_blocks(body, msg):
    """Indeed's plain-text alerts hide every job behind an opaque tracking link (no job key), so read
    the blocks instead: title / "Company - Location" / [salary] / snippet / age / link. A one-job
    "Title @ Company" recommendation email lists title, company and location above "View job:".
    The tracking links are never opened; the url stays empty and Claude finds the public posting."""
    import hashlib
    out = []

    def add(title, employer, location, context, link):
        key = hashlib.sha1(f"{title}|{employer}".lower().encode()).hexdigest()[:10]
        out.append({"id": f"IND-{key}", "url": "", "email_link": link, "title": title, "employer": employer,
                    "location": location, "posted": msg["date"], "text": "", "context": context[:400],
                    "source": "IND", "alert": msg["subject"]})

    for block in re.split(r"\n\s*\n", body):
        ls = [re.sub(r"\s+", " ", l).strip() for l in block.splitlines() if l.strip()]
        if len(ls) >= 4 and ls[-1].startswith("http") and AGE.match(ls[-2]) and " - " in ls[1]:
            employer, location = ls[1].rsplit(" - ", 1)
            add(ls[0], employer, location, " | ".join(ls[:-1]), ls[-1])
    subj = msg.get("subject", "")
    view = re.search(r"^View job:\s*(https?://\S+)", body, re.M)
    if not out and " @ " in subj and view:
        title, employer = subj.rsplit(" @ ", 1)
        ls = [re.sub(r"\s+", " ", l).strip() for l in body.splitlines() if l.strip()]
        i = next((k for k, l in enumerate(ls) if l == title.strip()), None)
        location = ls[i + 2] if i is not None and i + 2 < len(ls) and ls[i + 1] == employer.strip() else ""
        add(title.strip(), employer.strip(), location, " | ".join(ls[i:i + 6]) if i is not None else subj,
            view.group(1))
    return out


def alert_name(cand):
    """Where a posting was found, for the Leads sheet: "Indeed alert: economist" or the email subject."""
    s = cand.get("alert") or ""
    m = re.search(r"job alert for (.+?) jobs? in ", s, re.I)
    board = {"IND": "Indeed", "LI": "LinkedIn"}.get(cand.get("source"), cand.get("source") or "")
    return f"{board} alert: {m.group(1)}" if m else f"{board} email: {s}" if s else board


def alert_jobs(msg):
    """Postings listed in one alert email: each link to a job page of that board, its link text
    (HTML part) or the title / company / location lines printed above it (plain-text part), and a
    few lines of context for Claude to check."""
    b = msg["board"]
    mm = email.message_from_string(msg["source"], policy=email.policy.default)
    part = mm.get_body(preferencelist=("plain", "html"))
    body = part.get_content() if part else ""
    anchor = {}
    if part is not None and part.get_content_type() == "text/html":
        def a_(x):
            u, t = html.unescape(x.group(1)), clean(x.group(2))
            if t:
                anchor.setdefault(norm_link(unwrap(u, b["link"])), t)
            return "\n" + t + "\n" + u + "\n"
        body = re.sub(r'<a [^>]*href="([^"]+)"[^>]*>(.*?)</a>', a_, body, flags=re.S | re.I)
        body = re.sub(r"<(script|style).*?</\1>", " ", body, flags=re.S | re.I)
        body = html.unescape(re.sub(r"<[^>]+>", "\n", re.sub(r"<br\s*/?>", "\n", body)))
    lines = [re.sub(r"\s+", " ", l).strip() for l in body.splitlines()]
    lines = [l for l in lines if l]
    if b["code"] == "IND" and part is not None and part.get_content_type() == "text/plain":
        found = indeed_blocks(body, msg)
        if found:
            return found
    out, seen = [], set()
    skip = re.compile(r"(view job|see job|apply|easy apply|new|promoted|actively recruiting|save)\b", re.I)
    for i, l in enumerate(lines):
        for u in re.findall(r"https?://[^\s<>\"]+", l):
            u = unwrap(html.unescape(u).rstrip(">)].,"), b["link"])
            if not re.search(b["link"], u):
                continue
            link = norm_link(u)
            if link in seen:
                continue
            seen.add(link)
            above = [x for x in lines[max(0, i - 6):i] if not x.startswith("http") and len(x) < 120 and not skip.match(x)][-3:]
            if anchor.get(link):     # HTML: link text = title; company and location follow the link
                title = anchor[link]
                rest = [x for x in lines[i + 1:i + 5] if not x.startswith("http") and len(x) < 120
                        and not skip.match(x) and x != title][:2]
            else:                    # plain text: title, company, location above the link
                title = above[0] if len(above) == 3 else above[-1] if above else ""
                rest = [x for x in above if x != title]
            company, location = (rest + ["", ""])[:2]
            if " - " in company and not location:
                company, location = company.split(" - ", 1)
            context = " | ".join(x for x in lines[max(0, i - 5):i + 4] if not x.startswith("http"))[:400]
            out.append({"id": f"{b['code']}-{(re.findall(r'[0-9a-f]{6,}', link) or [link])[-1]}",
                        "url": "https://www." + link if b["code"] in ("LI", "IND") else re.sub(r"[?#].*$", "", u),
                        "title": title, "employer": company, "location": location, "posted": msg["date"],
                        "text": "", "context": context, "source": b["code"], "alert": msg["subject"]})
    return out


# ---------------- filter

def matcher(ind):
    cats = ind.get("categories") or {}
    excl = [w.lower() for w in ind.get("exclude_titles") or DEFAULT_EXCLUDE]

    def category(title, extra=(), alert=False):
        """Category name; None = drop. Alert emails were already filtered by the user's own alert
        keywords, so only the exclusions apply to them ("Other" when no category word matches)."""
        t = " " + title.lower() + " "
        if any(w in t for w in excl):
            return None
        for name, words in cats.items():
            if any(w.lower() in t for w in words):
                return name
        if any(w.lower() in t for w in extra):     # company-specific keywords (e.g. "Associate" at a consulting firm)
            return next(iter(cats), "Other")
        return "Other" if alert and title else None
    return category


def main():
    cfg = load_config()
    ind = cfg.get("industry") or {}
    out_dir = os.path.dirname(tracker_path(cfg, "industry"))
    os.makedirs(out_dir, exist_ok=True)
    state_path = os.path.join(out_dir, "industry_scan_state.json")
    state = json.load(open(state_path)) if os.path.exists(state_path) else {"seen": [], "mail_seen": []}
    seen = set() if "--all" in sys.argv else set(state["seen"])
    category = matcher(ind)

    tracked = set()
    for _, p in trackers(cfg):
        wb = load_workbook(p, read_only=True)
        for name in ("Tracker", "Leads"):
            if name in wb.sheetnames:
                for row in wb[name].iter_rows(values_only=True):
                    # links anywhere in a cell, including "also posted at <link>" in Flags
                    tracked |= {norm_link(u) for v in row if isinstance(v, str)
                                for u in re.findall(r"https?://[^\s,;]+", v)}

    posts, errors, stats = [], [], {}
    fns = {"greenhouse": greenhouse, "lever": lever, "ashby": ashby, "amazon": amazon}
    for c in ind.get("companies", []) if ind.get("scan_companies", True) or "--companies" in sys.argv else []:
        try:
            keep = lambda t: category(t, c.get("title_keywords", ()))
            got = ({"workday": workday, "google": google, "eightfold": eightfold}[c["ats"]](c, keep)
                   if c["ats"] in ("workday", "google", "eightfold") else fns[c["ats"]](c))
            for x in got:
                x.update({"source": c["ats"], "employer": c["name"], "type": c.get("type", ""),
                          "title_keywords": c.get("title_keywords", [])})
            posts += got
        except Exception as e:      # keep going if one site is down or changed
            errors.append(f"{c['name']} ({c['ats']}): {type(e).__name__}: {e}")
    m = ind.get("alert_mail") or {}
    mail_checked = []
    if m.get("account") and not m["account"].startswith("<") and "--no-mail" not in sys.argv:
        try:
            msgs = alert_messages(m, set(state.get("mail_seen", [])))
            mail_checked = [x["id"] for x in msgs]
            for msg in msgs:
                posts += alert_jobs(msg)
        except SystemExit as e:
            errors.append(f"mail: {e}")

    for i, a in enumerate(sys.argv):
        if a == "--board-file" and i + 1 < len(sys.argv):
            # postings read in the browser from a board that is all economics jobs (NABE EconJobs):
            # like alert emails, only the title exclusions apply
            for x in json.load(open(sys.argv[i + 1], encoding="utf-8")):
                posts.append({**x, "text": x.get("text", ""), "alert": "board", "posted": x.get("posted", "")})

    new = []
    for x in posts:
        if x["id"] in seen or norm_link(x["url"]) in tracked:
            stats["seen"] = stats.get("seen", 0) + 1
            continue
        cat = category(x["title"], x.pop("title_keywords", ()), alert=bool(x.get("alert")))
        if not cat:
            stats["title"] = stats.get("title", 0) + 1
            continue
        x["category"] = cat
        x["visa_text"] = " ".join(dict.fromkeys(s.strip() for s in VISA.findall(x["text"])))[:800]
        x["text"] = x["text"][:6000]
        x["summary"] = x["text"][:1200]
        new.append(x)

    # the same job from an alert and a career page: keep the career page (it has the full text)
    def key(x):
        emp = re.sub(r"[^a-z]", "", re.sub(r"\b(inc|llc|the|group)\b", "", x["employer"].lower()))[:20]
        return emp + "|" + " ".join(re.findall(r"[a-z]+", x["title"].lower())[:5])
    merged = {}
    for x in sorted(new, key=lambda x: bool(x.get("alert"))):
        k = key(x)
        if k in merged:
            if x["url"]:
                merged[k].setdefault("also_at", []).append(x["url"])
            stats["duplicate"] = stats.get("duplicate", 0) + 1
        else:
            merged[k] = x
    new = sorted(merged.values(), key=lambda x: (x["category"], x["employer"]))

    today = dt.date.today().isoformat()
    out_path = os.path.join(out_dir, "industry_scan_new.json")
    json.dump({"date": today, "errors": errors, "filtered_out": stats, "candidates": new},
              open(out_path, "w"), indent=1, ensure_ascii=False)
    state["seen"] = sorted(set(state["seen"]) | {x["id"] for x in posts})
    state["mail_seen"] = sorted(set(state.get("mail_seen", [])) | set(mail_checked))
    state["last_run"] = today
    json.dump(state, open(state_path, "w"))
    print(f"{len(posts)} postings fetched ({len(mail_checked)} alert emails), {len(new)} new candidates -> {out_path}")
    print("filtered out:", stats, "| errors:", errors or "none")


def view():
    cfg = load_config()
    data = json.load(open(os.path.join(os.path.dirname(tracker_path(cfg, "industry")), "industry_scan_new.json"),
                          encoding="utf-8"))
    c = data["candidates"]
    if sys.argv[1] == "show":
        start = int(sys.argv[2]) if len(sys.argv) > 2 else 0
        count = int(sys.argv[3]) if len(sys.argv) > 3 else 40
        if start == 0:
            print(f"scan date {data['date']} | {len(c)} candidates | errors: {data['errors'] or 'none'} "
                  f"| filtered out: {data['filtered_out']}\n")
        for n in range(start, min(start + count, len(c))):
            x = c[n]
            print(f"#{n} [{x['source']}] {x['category']} | {x['employer']} | {x['title']} | {x['location']} | posted {x.get('posted') or '?'}")
            if x.get("context"):
                print(f"   EMAIL: {x['context']}")
            if x.get("visa_text"):
                print(f"   VISA: {x['visa_text'][:300]}")
            print(f"   {x['summary'][:500] or '(alert email: no text; open the link)'}\n")
        if start + count < len(c):
            print(f"... next: show {start + count} {count}")
    else:
        for n in sys.argv[2:]:
            x = c[int(n)]
            print(f"#{n} [{x['source']}] {x['employer']} | {x['title']} | {x['url']}")
            print(f"also at: {x.get('also_at') or '-'} | {x['location']} | category {x['category']}")
            print(f"VISA: {x.get('visa_text') or '(nothing in ad)'}")
            if x.get("context"):
                print(f"EMAIL: {x['context']}")
            print((x["text"] or "(alert email: no text; open the link)") + "\n" + "-" * 60)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] in ("show", "text"):
        view()
    else:
        main()
