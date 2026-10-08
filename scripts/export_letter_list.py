"""Export a shareable letter-request list from the tracker.

  python3 scripts/export_letter_list.py

Reads the Tracker sheet, keeps rows with Letters? = Yes (skipping Withdrawn/Rejected), sorts them by deadline, and writes a
separate, read-only-style workbook to config["letter_share_file"] (e.g. a file in OneDrive or
Google Drive). Share that file's link once with your letter writers; re-running this script
overwrites the same file, so the link keeps working and always shows the latest list.
Only the columns writers need are copied (deadline, when letters are due, type of job, where to
submit, link, your status and the date you applied), plus one column per writer; notes stay private.
Due is the tracker's Letters Due, or the Deadline when that is blank.

After those come the hand-filled columns: each writer's status (Sent / Waiting), then the comments,
"<your first name> Comments" and "<writer> Comments" for each writer, filled in by hand in the
shared Google Sheet. Each refresh reads
these cells back first and keeps them, matching positions by link (or employer and position);
without a Google Sheet they are kept from the previous Excel file. (The tracker's own writer
columns are separate: the mail check marks Received there.)

If config["letter_share_gsheet"] is set, the list is written into that Google Sheet (see
scripts/gsheet.py for the one-time setup); the Excel file is still saved as a backup.
"""
import datetime as dt
import glob
import json
import os
import sys
import time

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from common import jm_path, load_config
from gsheet import WRITER_CHOICES, row_keys

# colors for the Status column, same as the tracker
STATUS_COLORS = [("Not started", "EDEDED", "595959"), ("In progress", "FFEB9C", "9C5700"),
                 ("Submitted", "C6EFCE", "006100"), ("Interview", "BDD7EE", "1F3864"),
                 ("Flyout", "9BC2E6", "1F3864"), ("Offer", "70AD47", "FFFFFF")]


def hand_columns(cfg):
    """[(header, width, kind, writer)] for the columns people fill in by hand; writer None = you."""
    me, ws = cfg["name"].split()[0], cfg["letter_writers"]
    return ([(w, 13, "status", w) for w in ws] + [(f"{me} Comments", 30, "comment", None)]
            + [(f"{w} Comments", 30, "comment", w) for w in ws])


def read_xlsx_status(path, headers):
    """{row key: {header: value}} for the given headers, from the previous Excel export."""
    if not os.path.exists(path):
        return {}
    sh = load_workbook(path)["Letter Requests"]
    col = {sh.cell(row=4, column=c).value: c for c in range(1, sh.max_column + 1)}
    if "Employer" not in col:
        return {}
    out = {}
    for r in range(5, sh.max_row + 1):
        got = {h: str(sh.cell(row=r, column=col[h]).value).strip() for h in headers
               if h in col and sh.cell(row=r, column=col[h]).value}
        if got:
            for k in row_keys(sh.cell(row=r, column=col["Link"]).value if "Link" in col else None,
                              sh.cell(row=r, column=col["Employer"]).value,
                              sh.cell(row=r, column=col["Position"]).value if "Position" in col else None):
                out[k] = got
    return out


def _state_path(cfg):
    return os.path.join(os.path.dirname(jm_path(cfg, cfg["tracker_file"])), "letter_share_state.json")


def recently_edited(cfg, dst, minutes=10):
    """True if someone other than this script saved the shared file in the last few minutes."""
    if not os.path.exists(dst):
        return False
    mtime = os.path.getmtime(dst)
    try:
        ours = json.load(open(_state_path(cfg))).get("mtime")
    except (OSError, ValueError):
        ours = None
    return (ours is None or abs(mtime - ours) > 0.001) and time.time() - mtime < minutes * 60


def remember_write(cfg, dst):
    if os.path.exists(dst):
        json.dump({"mtime": os.path.getmtime(dst)}, open(_state_path(cfg), "w"))


def conflicted_copies(dst):
    stem, ext = os.path.splitext(os.path.basename(dst))
    return [p for p in glob.glob(os.path.join(os.path.dirname(dst), glob.escape(stem) + "*" + ext))
            if p != dst and ("conflict" in p.lower() or "冲突" in p)]


def export():
    cfg = load_config()
    src = jm_path(cfg, cfg["tracker_file"])
    dst = os.path.expanduser(cfg["letter_share_file"])
    hand = [h for h, *_ in hand_columns(cfg)]   # the hand-filled headers to keep
    use_gsheet = bool(cfg.get("letter_share_gsheet"))
    prev, gsheet_ok = None, False
    if use_gsheet:
        try:
            from gsheet import read_writer_status
            prev, gsheet_ok = read_writer_status(cfg, hand), True
        except Exception as e:
            # never overwrite the sheet without having read what the writers entered
            print(f"WARNING: could not read the Google Sheet ({type(e).__name__}: {e}); it is left as is.")
    if prev is None:
        if recently_edited(cfg, dst) and "--force" not in sys.argv:
            print("Skipped: the shared list was edited in the last 10 minutes (someone may still be typing). "
                  "It will refresh on the next run; use --force to refresh now.")
            return
        prev = read_xlsx_status(dst, hand)
        for copy in conflicted_copies(dst):
            # Dropbox keeps both versions when two people save at once; take what was typed there too
            for k, got in read_xlsx_status(copy, hand).items():
                prev[k] = {**got, **prev.get(k, {})}
            print(f"NOTE: found {os.path.basename(copy)}; its entries were merged in. You can delete it.")
    _write(cfg, src, dst, hand, prev, gsheet_ok)
    remember_write(cfg, dst)


def _write(cfg, src, dst, hand, prev, gsheet_ok):
    ws = load_workbook(src, data_only=False)["Tracker"]
    hdr = {ws.cell(row=4, column=c).value: c for c in range(1, ws.max_column + 1) if ws.cell(row=4, column=c).value}
    need = ["Employer", "Position", "Type", "Apply Via", "Link", "Deadline", "Status", "Letters?"]
    missing = [h for h in need if h not in hdr]
    if missing:
        raise SystemExit(f"Tracker is missing columns: {missing}")

    rows = []
    for r in range(5, ws.max_row + 1):
        get = lambda h: ws.cell(row=r, column=hdr[h]).value
        emp = get("Employer")
        if not emp or str(emp).startswith("EXAMPLE") or str(get("Letters?")).strip() != "Yes":
            continue
        status = str(get("Status") or "").strip() or "Not started"
        if status in ("Withdrawn", "Rejected"):
            continue
        d = get("Deadline")
        if isinstance(d, dt.datetime):
            d = d.date()
        due = (get("Letters Due") if "Letters Due" in hdr else None) or d
        if isinstance(due, dt.datetime):
            due = due.date()
        sub = get("Submitted") if "Submitted" in hdr else None
        if isinstance(sub, dt.datetime):
            sub = sub.date()
        got = dict(next((prev[k] for k in row_keys(get("Link"), emp, get("Position")) if k in prev), {}))
        for w in cfg["letter_writers"]:       # the mail check confirmed this letter: show it to everyone
            if w in hdr and str(get(w) or "").strip() in ("Received", "Uploaded"):
                got[w] = "Received"
        rows.append({"Deadline": d, "Due": due, "Submitted": sub, "Employer": emp, "Position": get("Position"), "Type": get("Type"),
                     "Apply Via": get("Apply Via"), "Link": get("Link"), "Status": status,
                     **{h: got.get(h, "") for h in hand}})
    rows.sort(key=lambda x: x["Deadline"] or dt.date(2099, 12, 31))

    FONT = "Arial"
    accent = cfg.get("accent_color", "0021A5")
    fb = Font(name=FONT, size=10)
    fh = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    side = Side(style="thin", color="D0D0D0")
    bd = Border(left=side, right=side, top=side, bottom=side)
    wrap = Alignment(wrap_text=True, vertical="top")
    ctr = Alignment(horizontal="center", vertical="top", wrap_text=True)

    wb = Workbook()
    sh = wb.active
    sh.title = "Letter Requests"
    sh["A1"] = f"Recommendation Letter Requests · {cfg['name']} · {cfg.get('season', '').replace('-', '–')}"
    sh["A1"].font = Font(name=FONT, size=15, bold=True, color=accent)
    sh["A2"] = (f"Positions that need a letter, soonest deadline first. Once I submit (Status = Submitted, "
                f"'I Applied On' filled), the application system sends each of you the upload request. "
                f"Due is when your letter should be in (usually when review begins). "
                f"Please mark your own column Sent / Waiting; it changes to Received when the system confirms your letter. "
                f"Comments are welcome in your Comments column; everything else is locked. Last updated "
                f"{dt.date.today().strftime('%B %-d, %Y')}. Thank you for your support!")
    sh["A2"].font = Font(name=FONT, size=10, italic=True, color="555555")
    # no formulas: the file is mostly viewed in a browser preview (Dropbox/OneDrive), which does not
    # recalculate them, so a "Days Left" formula would show up blank
    cols = [("#", 5), ("Deadline", 13), ("Due", 13), ("Employer", 30), ("Position", 32), ("Type", 16),
            ("Submit Letter Via", 20), ("Link", 30), ("Status", 13), ("I Applied On", 13)]
    LINK, STATUS, DATES = 8, 9, (2, 3, 10)
    HAND = [(len(cols) + 1 + i, kind, who) for i, (_, _, kind, who) in enumerate(hand_columns(cfg))]
    cols += [(h, w) for h, w, *_ in hand_columns(cfg)]
    WCOLS = [c for c, kind, _ in HAND if kind == "status"]
    for i, (h, w) in enumerate(cols, start=1):
        c = sh.cell(row=4, column=i, value=h)
        c.font = fh; c.fill = PatternFill("solid", fgColor=accent); c.alignment = ctr; c.border = bd
        sh.column_dimensions[L(i)].width = w
    table = [[k, x["Deadline"], x["Due"], x["Employer"], x["Position"], x["Type"], x["Apply Via"], x["Link"], x["Status"], x["Submitted"]]
             + [x[h] for h in hand] for k, x in enumerate(rows, start=1)]
    for k, (x, vals) in enumerate(zip(rows, table), start=1):
        r = 4 + k
        for i, v in enumerate(vals, start=1):
            c = sh.cell(row=r, column=i, value=v)
            c.font = fb; c.border = bd
            c.alignment = ctr if i in (1, STATUS, *DATES, *WCOLS) else wrap
        for i in DATES:
            sh.cell(row=r, column=i).number_format = "mmm d, yyyy"
        if x["Link"]:
            sh.cell(row=r, column=LINK).hyperlink = x["Link"]
            sh.cell(row=r, column=LINK).font = Font(name=FONT, size=10, color="0563C1", underline="single")
    if not rows:
        sh["A5"] = "No positions need letters yet."
        sh["A5"].font = fb
    last = 4 + max(len(rows), 1)
    for c0, c1, choices in [(STATUS, STATUS, STATUS_COLORS)] + [(c, c, WRITER_CHOICES) for c in WCOLS]:
        a, b = L(c0), L(c1)
        for text, bg, fg in choices:
            sh.conditional_formatting.add(f"{a}5:{b}{last}", FormulaRule(
                formula=[f'{a}5="{text}"'], fill=PatternFill("solid", fgColor=bg), font=Font(name=FONT, color=fg)))
    sh.freeze_panes = "A5"
    sh.sheet_view.showGridLines = False
    # locked except the hand-filled columns (writer status and comments), which anyone with the link can edit
    for c, kind, _ in HAND:
        for r in range(5, max(last, 5) + 1):
            sh.cell(row=r, column=c).protection = Protection(locked=False)
    if rows:
        dv = DataValidation(type="list", formula1='"' + ",".join(t for t, *_ in WRITER_CHOICES) + '"', allow_blank=True)
        sh.add_data_validation(dv)
        for c in WCOLS:
            dv.add(f"{L(c)}5:{L(c)}{last}")
    sh.protection.sheet = True          # no password, so you can still unprotect locally
    # ...but still let viewers hide/resize columns and rows, sort and filter (False = allowed)
    for opt in ("formatColumns", "formatRows", "sort", "autoFilter"):
        setattr(sh.protection, opt, False)
    # columns to hide on every export (headers, e.g. ["Submit Letter Via"]); set in config.json
    for i, (h, _) in enumerate(cols, start=1):
        if h in cfg.get("letter_share_hide", []):
            sh.column_dimensions[L(i)].hidden = True
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    wb.save(dst)
    print(f"{len(rows)} position(s) -> {dst}")

    if gsheet_ok:
        # a Google failure (offline, bad key) must not stop the mail check / tracker update that called us
        try:
            from gsheet import push_letter_list
            url = push_letter_list(cfg, sh["A1"].value, sh["A2"].value, cols, table, link_col=LINK, date_cols=DATES,
                                   status_col=STATUS, status_colors=STATUS_COLORS, hand_cols=HAND,
                                   hide=cfg.get("letter_share_hide", []))
            print(f"{len(rows)} position(s) -> Google Sheet {url}")
        except Exception as e:
            print(f"WARNING: Google Sheet not updated ({type(e).__name__}: {e}); the Excel file was saved.")


if __name__ == "__main__":
    export()
