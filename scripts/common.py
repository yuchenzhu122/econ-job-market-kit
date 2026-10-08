"""Shared helpers: load config.json from the repo root and resolve paths."""
import json
import os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def load_config():
    path = os.path.join(ROOT, "config.json")
    if not os.path.exists(path):
        raise SystemExit("config.json not found. Copy config.example.json to config.json and fill it in.")
    with open(path, encoding="utf-8") as f:
        cfg = json.load(f)
    cfg["job_market_dir"] = os.path.expanduser(cfg["job_market_dir"])
    return cfg


def jm_path(cfg, rel):
    """Absolute path inside the job market folder."""
    return os.path.join(cfg["job_market_dir"], rel)


def tex_escape(s):
    import re
    for a, b in (("&", r"\&"), ("%", r"\%"), ("$", r"\$"), ("#", r"\#"), ("_", r"\_")):
        s = s.replace(a, b)
    s = s.replace("’", "'").replace("“", "``").replace("”", "''")
    return re.sub(r'"(\S)', r"``\1", s).replace('"', "''")


def run_pdflatex(workdir, stem):
    import subprocess
    for _ in range(2):
        subprocess.run(["pdflatex", "-interaction=nonstopmode", stem + ".tex"], cwd=workdir,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    for ext in (".aux", ".log", ".out"):
        p = os.path.join(workdir, stem + ext)
        if os.path.exists(p):
            os.remove(p)
    return os.path.join(workdir, stem + ".pdf")


def letterhead(cfg):
    """LaTeX preamble + header shared by the cover letter and statements."""
    return r"""\documentclass[11pt]{article}
\usepackage[letterpaper,margin=1in]{geometry}
\usepackage[T1]{fontenc}
\usepackage{mathptmx}
\usepackage[dvipsnames]{xcolor}
\usepackage[hidelinks]{hyperref}
\definecolor{accent}{HTML}{""" + cfg.get("accent_color", "0021A5") + r"""}
\pagestyle{empty}
\setlength{\parindent}{0pt}
"""


def restore_dropdowns(wb, writers):
    """Excel re-saves the Tracker's cross-sheet dropdowns in an extension format openpyxl drops on
    load. Call before saving the workbook with openpyxl to put the dropdowns back."""
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.datavalidation import DataValidation
    tr, li = wb["Tracker"], wb["Lists"]
    have = {str(c) for v in tr.data_validations.dataValidation for c in v.sqref.ranges}
    lists = {li.cell(row=3, column=c).value: L(c) for c in range(1, li.max_column + 1) if li.cell(row=3, column=c).value}
    H = {tr.cell(row=4, column=c).value: L(c) for c in range(1, tr.max_column + 1) if tr.cell(row=4, column=c).value}
    want = {"Track": "Track", "Type": "Type", "Apply Via": "Platform", "Status": "Status",
            "Letters?": "Yes/No", "Cover Letter": "Cover Letter", **{w: "Letter" for w in writers}}
    last = max(154, tr.max_row)
    for head, name in want.items():
        if head not in H or name not in lists:
            continue
        rng = f"{H[head]}5:{H[head]}{last}"
        if any(r.startswith(H[head] + "5:") for r in have):
            continue
        col = lists[name]
        v = DataValidation(type="list", formula1=f"=Lists!${col}$4:${col}$30", allow_blank=True)
        tr.add_data_validation(v)
        v.add(rng)


def excel_open_names():
    import subprocess
    if subprocess.run(["pgrep", "-x", "Microsoft Excel"], capture_output=True).returncode != 0:
        return ""
    r = subprocess.run(["osascript", "-e", 'tell application "Microsoft Excel" to get name of workbooks'],
                       capture_output=True, text=True)
    return r.stdout


def close_in_excel(path):
    """If Excel has the workbook open, save and close it. Returns True if it was open (reopen later
    with reopen_in_excel). Asks Excel directly; the ~$ lock file is unreliable."""
    import subprocess, time
    name = os.path.basename(path)
    was_open = name in excel_open_names()
    if was_open:
        r = subprocess.run(["osascript", "-e", f'tell application "Microsoft Excel" to close workbook "{name}" saving yes'],
                           capture_output=True, text=True)
        if r.returncode != 0:
            raise SystemExit(f"Could not ask Excel to save and close the tracker ({r.stderr.strip()}). "
                             "Save and close it yourself, then run this again.")
        for _ in range(40):
            if name not in excel_open_names():
                break
            time.sleep(0.25)
        else:
            raise SystemExit("Excel still has the tracker open. Save and close it, then run this again.")
    lock = os.path.join(os.path.dirname(path), "~$" + name)
    if os.path.exists(lock):
        try:
            os.remove(lock)
        except OSError:
            pass
    return was_open


def reopen_in_excel(path, was_open):
    import subprocess
    if was_open:
        subprocess.run(["open", path])


def tracker_headers(ws):
    return {ws.cell(row=4, column=c).value: c for c in range(1, ws.max_column + 1) if ws.cell(row=4, column=c).value}


def ensure_submitted_column(ws):
    """Add a 'Submitted' date column at the right edge of an older Tracker sheet. Returns its index."""
    return ensure_date_column(ws, "Submitted")


def ensure_date_column(ws, name):
    """Add a date column (e.g. 'Submitted', 'Letters Due') at the right edge of an older Tracker
    sheet if it is missing. Returns its index."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.datavalidation import DataValidation
    H = tracker_headers(ws)
    if name in H:
        return H[name]
    col = max(H.values()) + 1
    src = ws.cell(row=4, column=H["Deadline"])
    c = ws.cell(row=4, column=col, value=name)
    c.font = Font(name=src.font.name, size=src.font.sz, bold=True, color=src.font.color.rgb if src.font.color else None)
    c.fill = PatternFill("solid", fgColor=src.fill.fgColor.rgb)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[L(col)].width = 12
    for name in ("_sort", "_order"):
        if name in H:
            ws.column_dimensions[L(H[name])].hidden = True
    for r in range(5, 155):
        ws.cell(row=r, column=col).number_format = "mmm d, yyyy"
    d = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True)
    d.error = "Enter a date, e.g. 11/15/2026."
    ws.add_data_validation(d)
    d.add(f"{L(col)}5:{L(col)}154")
    return col
