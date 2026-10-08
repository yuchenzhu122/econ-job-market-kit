import datetime as dt
import os
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import jm_path, load_config

CFG = load_config()
OUT = jm_path(CFG, CFG["tracker_file"])
if os.path.exists(OUT) and "--force" not in sys.argv:
    raise SystemExit(f"{OUT} already exists. Re-run with --force to overwrite it (all rows will be lost).")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
SEASON = CFG.get("season", "2026-27").replace("-", "–")
FONT = "Arial"; NAVY = CFG.get("accent_color", "0021A5")
fb = Font(name=FONT, size=10); fbold = Font(name=FONT, size=10, bold=True)
fh = Font(name=FONT, size=10, bold=True, color="FFFFFF")
ft = Font(name=FONT, size=16, bold=True, color=NAVY)
fsub = Font(name=FONT, size=10, italic=True, color="555555")
fex = Font(name=FONT, size=10, italic=True, color="7F7F7F")
flink = Font(name=FONT, size=10, color="0563C1", underline="single")
hfill = PatternFill("solid", fgColor=NAVY); grey = PatternFill("solid", fgColor="F2F2F2")
side = Side(style="thin", color="D0D0D0"); bd = Border(left=side, right=side, top=side, bottom=side)
wrap = Alignment(wrap_text=True, vertical="top")
ctr = Alignment(horizontal="center", vertical="top", wrap_text=True)

WRITERS = CFG["letter_writers"]
N = 150; FIRST = 5; LAST = FIRST + N - 1

def cfr(ws, rng, formula, fill, color=None, bold=False):
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[formula], fill=PatternFill("solid", fgColor=fill),
        font=Font(name=FONT, color=color, bold=bold) if (color or bold) else None))

wb = Workbook()

# ---------------- Lists
ls = wb.active; ls.title = "Lists"
lists = {
    "Track": ["Academic", "Industry", "Government / Policy"],
    "Type": ["Tenure-track", "Teaching-focused", "Postdoc", "Visiting", "Fed / Central bank",
             "Government", "Think tank / Research", "Tech", "Consulting", "Finance", "Other"],
    "Platform": ["EconJobMarket", "JOE (AEA)", "AcademicJobsOnline", "Interfolio",
                 "HigherEdJobs", "Chronicle Jobs", "Inside Higher Ed Careers", "Employer website", "Email", "LinkedIn / Handshake", "USAJOBS", "Other"],
    "Status": ["Not started", "In progress", "Submitted", "Interview", "Flyout", "Offer",
               "Rejected", "Withdrawn"],
    "Yes/No": ["Yes", "No"],
    "Letter": ["Received"],
    "Cover Letter": ["To write", "Drafted", "Final", "Not needed"],
}
ls["A1"] = "Dropdown options — add new options at the bottom of any column."; ls["A1"].font = fbold
for c, (name, vals) in enumerate(lists.items(), start=1):
    h = ls.cell(row=3, column=c, value=name); h.font = fh; h.fill = hfill
    for r, v in enumerate(vals, start=4):
        ls.cell(row=r, column=c, value=v).font = fb
    ls.column_dimensions[L(c)].width = 24
RANGE = {n: f"Lists!${L(c)}$4:${L(c)}$30" for c, n in enumerate(lists, start=1)}

# ---------------- Tracker
tr = wb.create_sheet("Tracker", 0)
cols = [  # header, width, kind, note
    ("Track", 13, "in", "Academic / Industry / Government / Policy"),
    ("Employer", 28, "in", None),
    ("Position", 28, "in", None),
    ("Type", 16, "in", "Optional"),
    ("Apply Via", 18, "in", "Where you submit (letter writers also upload here)"),
    ("Link", 24, "link", None),
    ("Deadline", 12, "date", "The date you apply by: the earliest date in the ad (review begins, priority date or deadline)."),
    ("Letters Due", 12, "date", "When letters should be in (also on the writers' list): the later 'full consideration' date if the ad gives two dates, the review date if it gives only that. Blank until known."),
    ("Days Left", 9, "auto", "Automatic. Red = 7 days or less. Disappears once Status is Submitted."),
    ("Status", 13, "in", None),
    ("Submitted", 12, "date", "Date you submitted. Type it yourself, or let the optional mail check fill it from confirmation emails."),
    ("Still Need", 26, "in", "Type what's missing, e.g. 'cover letter, teaching evals'. Leave blank when everything is ready."),
    ("Cover Letter", 11, "in", "To write / Drafted / Final / Not needed. Ask Claude: 'write the cover letter for <Employer>' and it updates this cell. Files are in 09_Cover_Letters."),
    ("Letters?", 9, "in", "Yes = this job needs recommendation letters (it will show up on the letter writers' list)."),
] + [(w, 12, "letter", "Received once this writer's letter has arrived. Set it yourself, or let the optional mail check fill it.")
     for w in WRITERS] + [
    ("Notes", 34, "in", "Anything else: interview dates, contacts, what to stress in the cover letter."),
    ("_sort", 6, "auto", None), ("_order", 6, "auto", None),
]
C = {h: L(i) for i, (h, *_r) in enumerate(cols, start=1)}
tr["A1"] = "Job Market Tracker · " + SEASON; tr["A1"].font = ft
tr["A2"] = ("One row per job. Shortcut: paste only the Link and ask Claude to fill the rest. Otherwise must-fill: Track, Employer, Position, Apply Via, Link, Deadline, Status, Letters?. "
            "Grey cells are automatic. Delete the two grey EXAMPLE rows first.")
tr["A2"].font = fsub
for i, (h, w, kind, note) in enumerate(cols, start=1):
    c = tr.cell(row=4, column=i, value=h); c.font = fh; c.fill = hfill; c.alignment = ctr; c.border = bd
    tr.column_dimensions[L(i)].width = w
    if note:
        c.comment = Comment(note, "Tracker")
tr.row_dimensions[4].height = 22

def dv(listname, col):
    v = DataValidation(type="list", formula1=f"={RANGE[listname]}", allow_blank=True)
    tr.add_data_validation(v); v.add(f"{col}{FIRST}:{col}{LAST}")
dv("Track", C["Track"]); dv("Type", C["Type"]); dv("Platform", C["Apply Via"])
dv("Status", C["Status"]); dv("Yes/No", C["Letters?"]); dv("Cover Letter", C["Cover Letter"])
for w in WRITERS:
    dv("Letter", C[w])
d = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True)
d.error = "Enter a date, e.g. 11/15/2026."; tr.add_data_validation(d); d.add(f"{C['Deadline']}{FIRST}:{C['Deadline']}{LAST}")
d.add(f"{C['Submitted']}{FIRST}:{C['Submitted']}{LAST}")
d.add(f"{C['Letters Due']}{FIRST}:{C['Letters Due']}{LAST}")

E, DL, ST, LT, SD = C["Employer"], C["Deadline"], C["Status"], C["Letters?"], C["_sort"]
for r in range(FIRST, LAST + 1):
    tr[f"{C['Days Left']}{r}"] = (f'=IF(OR({DL}{r}="",{ST}{r}="Submitted",{ST}{r}="Interview",{ST}{r}="Flyout",'
                                   f'{ST}{r}="Offer",{ST}{r}="Rejected",{ST}{r}="Withdrawn"),"",{DL}{r}-TODAY())')
    tr[f"{SD}{r}"] = f'=IF({DL}{r}="",DATE(2099,12,31),{DL}{r})'
    tr[f"{C['_order']}{r}"] = (f'=IF(AND({E}{r}<>"",{LT}{r}="Yes"),'
                               f'COUNTIFS(${LT}${FIRST}:${LT}${LAST},"Yes",${SD}${FIRST}:${SD}${LAST},"<"&{SD}{r})'
                               f'+COUNTIFS(${LT}${FIRST}:{LT}{r},"Yes",${SD}${FIRST}:{SD}{r},{SD}{r}),"")')
    for i, (h, w, kind, note) in enumerate(cols, start=1):
        cell = tr.cell(row=r, column=i); cell.font = fb; cell.border = bd
        cell.alignment = ctr if kind in ("date", "auto", "letter") or h in ("Letters?", "Status", "Track") else wrap
        if kind == "auto":
            cell.fill = grey
        if kind == "date":
            cell.number_format = "mmm d, yyyy"
    tr[f"{C['Days Left']}{r}"].number_format = '0;"late "0;"today"'
for h in ("_sort", "_order"):
    tr.column_dimensions[C[h]].hidden = True

examples = [
    {"Track": "Academic", "Employer": "EXAMPLE – State University", "Position": "Assistant Professor (Labor/Education)",
     "Type": "Tenure-track", "Apply Via": "EconJobMarket", "Link": "https://www.aeaweb.org/joe/",
     "Deadline": dt.date(2026, 11, 15), "Status": "In progress", "Still Need": "teaching evals", "Cover Letter": "Drafted",
     "Letters?": "Yes", WRITERS[0]: "Received",
     "Notes": "Mention their education policy center in the cover letter."},
    {"Track": "Industry", "Employer": "EXAMPLE – Tech Company", "Position": "Economist",
     "Type": "Tech", "Apply Via": "Employer website", "Link": "https://example.com/careers",
     "Deadline": dt.date(2026, 12, 1), "Status": "Not started", "Still Need": "1-page resume", "Cover Letter": "Not needed",
     "Letters?": "No", "Notes": "Ask alum on the team for a referral."},
]
for k, ex in enumerate(examples):
    r = FIRST + k
    for h, v in ex.items():
        c = tr[f"{C[h]}{r}"]; c.value = v; c.font = fex
        if h == "Link":
            c.hyperlink = v
tr.freeze_panes = f"C{FIRST}"
tr.auto_filter.ref = f"A4:{C['Notes']}{LAST}"

rk = f"{C['Track']}{FIRST}:{C['Track']}{LAST}"; tk = f"${C['Track']}{FIRST}"
cfr(tr, rk, f'{tk}="Academic"', "DDEBF7", "1F3864", True)
cfr(tr, rk, f'{tk}="Industry"', "E2EFDA", "375623", True)
cfr(tr, rk, f'{tk}="Government / Policy"', "FFF2CC", "7F6000", True)
dk = f"${C['Days Left']}{FIRST}"; drng = f"{C['Days Left']}{FIRST}:{C['Days Left']}{LAST}"
cfr(tr, drng, f"AND(ISNUMBER({dk}),{dk}<0)", "C00000", "FFFFFF", True)
cfr(tr, drng, f"AND(ISNUMBER({dk}),{dk}<=7)", "F8CBAD", "9C0006", True)
cfr(tr, drng, f"AND(ISNUMBER({dk}),{dk}<=14)", "FFEB9C", "9C5700")
sk = f"${ST}{FIRST}"; srng = f"{ST}{FIRST}:{ST}{LAST}"
for val, f1, c1 in [("Submitted", "C6EFCE", "006100"), ("Interview", "BDD7EE", "1F3864"), ("Flyout", "9BC2E6", "1F3864"),
                    ("Offer", "70AD47", "FFFFFF"), ("In progress", "FFEB9C", "9C5700"),
                    ("Rejected", "EDEDED", "7F7F7F"), ("Withdrawn", "EDEDED", "7F7F7F")]:
    cfr(tr, srng, f'{sk}="{val}"', f1, c1)
if WRITERS:
    wf, wl = C[WRITERS[0]], C[WRITERS[-1]]
    cfr(tr, f"{wf}{FIRST}:{wl}{LAST}", f'{wf}{FIRST}="Received"', "C6EFCE", "006100")
clk = f"${C['Cover Letter']}{FIRST}"; clr = f"{C['Cover Letter']}{FIRST}:{C['Cover Letter']}{LAST}"
cfr(tr, clr, f'{clk}="To write"', "F8CBAD", "9C0006")
cfr(tr, clr, f'{clk}="Drafted"', "FFEB9C", "9C5700")
cfr(tr, clr, f'{clk}="Final"', "C6EFCE", "006100")
sn = f"${C['Still Need']}{FIRST}"
cfr(tr, f"{C['Still Need']}{FIRST}:{C['Still Need']}{LAST}", f'AND({sn}<>"",${E}{FIRST}<>"")', "FFF2CC")

# ---------------- Summary
sm = wb.create_sheet("Summary", 2)
sm["A1"] = "Summary"; sm["A1"].font = ft
tracks = lists["Track"]; statuses = lists["Status"]
sm.cell(row=3, column=1, value="Status").font = fh; sm.cell(row=3, column=1).fill = hfill
for j, t in enumerate(tracks + ["Total"], start=2):
    c = sm.cell(row=3, column=j, value=t); c.font = fh; c.fill = hfill; c.alignment = ctr
trk = f"Tracker!${C['Track']}${FIRST}:${C['Track']}${LAST}"; sts = f"Tracker!${ST}${FIRST}:${ST}${LAST}"
for i, s in enumerate(statuses, start=4):
    sm.cell(row=i, column=1, value=s).font = fb
    for j, t in enumerate(tracks, start=2):
        sm.cell(row=i, column=j, value=f'=COUNTIFS({trk},"{t}",{sts},$A{i})').font = fb
    sm.cell(row=i, column=len(tracks) + 2, value=f"=SUM(B{i}:{L(len(tracks) + 1)}{i})").font = fbold
tot = 4 + len(statuses)
sm.cell(row=tot, column=1, value="Total").font = fbold
for j in range(2, len(tracks) + 3):
    sm.cell(row=tot, column=j, value=f"=SUM({L(j)}4:{L(j)}{tot - 1})").font = fbold
for r in range(3, tot + 1):
    for j in range(1, len(tracks) + 3):
        sm.cell(row=r, column=j).border = bd
dlc = f"Tracker!${C['Days Left']}${FIRST}:${C['Days Left']}${LAST}"
r0 = tot + 2
for k, (lab, f) in enumerate([("Due within 7 days (not submitted)", f'=COUNTIFS({dlc},">=0",{dlc},"<=7")'),
                              ("Due within 14 days (not submitted)", f'=COUNTIFS({dlc},">=0",{dlc},"<=14")'),
                              ("Past deadline, not submitted", f'=COUNTIF({dlc},"<0")')], start=r0):
    sm.cell(row=k, column=1, value=lab).font = fb; sm.cell(row=k, column=2, value=f).font = fbold
    sm.cell(row=k, column=1).border = bd; sm.cell(row=k, column=2).border = bd
sm.column_dimensions["A"].width = 36
for j in range(2, 6):
    sm.column_dimensions[L(j)].width = 18

# ---------------- Key Dates
kd = wb.create_sheet("Key Dates", 3)
kd["A1"] = "Key dates · 2026–27 (AEA; update each season)"; kd["A1"].font = ft
for j, h in enumerate(["When", "What", "Source"], start=1):
    c = kd.cell(row=3, column=j, value=h); c.font = fh; c.fill = hfill
dates = [
    ("Oct–early Dec 2026", "Most academic application deadlines. Send letter writers your list early.", "Check each posting"),
    (dt.date(2026, 11, 10), "AEA signaling opens (pick up to 2 employers).", "https://www.aeaweb.org/joe/signal"),
    (dt.date(2026, 12, 1), "AEA signaling closes.", "https://www.aeaweb.org/joe/signal"),
    (dt.date(2026, 12, 3), "Signals sent to employers; interview invitations come after this.", "https://www.aeaweb.org/news/guidance-job-market-sept-9-2026"),
    ("Dec 2026 – Jan 2027", "First-round interviews, online.", "https://www.aeaweb.org/news/guidance-job-market-sept-9-2026"),
    ("Jan 3–5, 2027", "ASSA meetings (no interviews during the meetings).", "https://www.aeaweb.org/news/guidance-job-market-sept-9-2026"),
    (dt.date(2027, 1, 31), "Offers should stay open at least until this date; at least 2 weeks to decide.", "https://www.aeaweb.org/news/guidance-job-market-sept-9-2026"),
    ("March 2027", "AEA Job Market Scramble.", "https://www.aeaweb.org/joe/scramble"),
]
for k, (a, b, s) in enumerate(dates, start=4):
    c = kd.cell(row=k, column=1, value=a); c.font = fb
    if isinstance(a, dt.date):
        c.number_format = "mmm d, yyyy"
    kd.cell(row=k, column=2, value=b).font = fb
    sc = kd.cell(row=k, column=3, value=s)
    if s.startswith("http"):
        sc.hyperlink = s; sc.font = flink
    else:
        sc.font = fb
    for j in range(1, 4):
        kd.cell(row=k, column=j).alignment = wrap; kd.cell(row=k, column=j).border = bd
kd.cell(row=4 + len(dates) + 1, column=1,
        value="Letters on EconJobMarket: each writer uploads once and the letter is reused for every EJM job.").font = fsub
for j, w in enumerate([20, 70, 50], start=1):
    kd.column_dimensions[L(j)].width = w

# ---------------- Read Me
rm = wb.create_sheet("Read Me", 0)
rm["A1"] = "How to use (30 seconds)"; rm["A1"].font = ft
steps = [
    ("1  Find a job", "Fastest: paste just the posting URL into the Link column, then ask Claude '补全追踪表' (fill the tracker) — it reads each link and fills the rest. Or fill the row yourself: Track, Employer, Position, Apply Via, Link, Deadline, Status, Letters?."),
    ("2  Missing stuff", "Type it in 'Still Need' (e.g. 'teaching evals'). Clear it when ready."),
    ("   Cover letter", "Ask Claude: 'write the cover letter for <Employer>'. It drafts a tailored letter into 09_Cover_Letters and sets the Cover Letter column to Drafted."),
    ("3  Letters", "Set Letters? = Yes. When you submit, the application system invites all your writers; they mark Sent / Waiting on the shared Letter Requests list. Under their names here, pick Received once their letter has arrived (or let the optional mail check do it from confirmation emails)."),
    ("4  Submitted", "Change Status to Submitted and type the date in Submitted, then click ▶ Update Tracker to refresh the writers' list. The countdown disappears. (Optional: the mail check does this from confirmation emails.)"),
    ("Watch", "Days Left turns yellow at 14 days and red at 7 days."),
    ("Share", "This tracker is private. Writers get only the Letter Requests list (Google Sheet, Excel backup) made by export_letter_list.py; it refreshes after every update."),
    ("Start", "Delete the two grey EXAMPLE rows on Tracker."),
]
for k, (a, b) in enumerate(steps, start=3):
    x = rm.cell(row=k, column=1, value=a); x.font = fbold; x.alignment = wrap
    y = rm.cell(row=k, column=2, value=b); y.font = fb; y.alignment = wrap
rm.column_dimensions["A"].width = 18; rm.column_dimensions["B"].width = 110

wb._sheets = [wb["Read Me"], wb["Tracker"], wb["Summary"], wb["Key Dates"], wb["Lists"]]
wb["Tracker"].sheet_properties.tabColor = NAVY
wb.active = 1
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title == "Lists"
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print(OUT)
